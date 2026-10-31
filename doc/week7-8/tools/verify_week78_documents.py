"""Check DOCX content/source fingerprints and render private XPS previews to PNG only."""

import argparse
import hashlib
import json
import zipfile
from pathlib import Path

import fitz
from docx import Document
from lxml import etree
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[3]
WEEK = ROOT / 'doc/week7-8'
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--render-index', type=Path, required=True)
    args = parser.parse_args()
    render = json.loads(args.render_index.read_text(encoding='utf-8-sig'))
    assert not render['pdf_generated']
    provenance_path = WEEK / '文档编制记录/编制与来源.json'
    provenance = json.loads(provenance_path.read_text(encoding='utf-8'))
    for name, expected in provenance['inputs_sha256'].items():
        assert digest(ROOT / name) == expected, f'Changed input: {name}'
    for items in provenance['inventory'].values():
        for item in items:
            assert digest(ROOT / item['path']) == item['sha256'], item['path']
    facts = json.loads((WEEK / 'content/week78_confirmed_facts.json').read_text(encoding='utf-8'))
    preview = args.render_index.parent / 'pages'
    preview.mkdir(exist_ok=True)
    results = []
    for index, entry in enumerate(render['documents'], 1):
        path = Path(entry['source'])
        assert path.parent == WEEK and path.suffix == '.docx'
        assert digest(path) == entry['sha256'], f'Changed after Word preview: {path.name}'
        with zipfile.ZipFile(path) as archive:
            assert archive.testzip() is None
            xml = etree.fromstring(archive.read('word/document.xml'))
            assert not xml.xpath('//w:ins | //w:del', namespaces=NS)
            text = ''.join(xml.xpath('//w:t/text()', namespaces=NS))
        for member in facts['members']:
            assert member['name'] in text
        assert '小组审阅稿' in text
        assert '小组成员：xx' not in text and '//<' not in text
        assert text.count('xx') == 0
        required = (['规范制定的主要目的', '规范的适用范围', '3  规范'] if index == 1 else [
            '待分析开源软件概述', '软件质量分析方法及工具', '人工分析代码的质量情况',
            '工具分析代码的质量情况', '小组分工', '下一步的改进计划',
            '2026-10-23 19:00—19:30', '357', '【待确认】',
        ])
        for expected in required:
            assert expected in text, (path.name, expected)
        document = Document(path)
        assert len(document.tables) == entry['tables']
        assert all(table.rows[0]._tr.xpath('./w:trPr/w:tblHeader') for table in document.tables)
        book = fitz.open(entry['xps'])
        assert len(book) == entry['pages']
        pages = []
        out_of_page = []
        for page_number, page in enumerate(book, 1):
            image_path = preview / f'document-{index}-page-{page_number:02}.png'
            page.get_pixmap(dpi=110, alpha=False).save(image_path)
            for block in page.get_text('blocks'):
                if block[4].strip() and (block[0] < -1 or block[1] < -1
                                         or block[2] > page.rect.width + 1
                                         or block[3] > page.rect.height + 1):
                    out_of_page.append({'page': page_number, 'bounds': block[:4]})
            pages.append({'page': page_number, 'characters': len(page.get_text()),
                          'image': str(image_path)})
        assert not out_of_page, out_of_page
        sheets = []
        for start in range(0, len(pages), 4):
            sheet = Image.new('RGB', (1020, 1480), '#D6DCE5')
            draw = ImageDraw.Draw(sheet)
            for offset, page in enumerate(pages[start:start + 4]):
                image = Image.open(page['image']).convert('RGB')
                image.thumbnail((480, 690))
                x = 20 + (offset % 2) * 505
                y = 30 + (offset // 2) * 730
                draw.text((x, y - 20), f'Document {index} / Page {page["page"]}', fill='black')
                sheet.paste(image, (x, y))
            sheet_path = preview / f'document-{index}-sheet-{start // 4 + 1}.png'
            sheet.save(sheet_path)
            sheets.append(str(sheet_path))
        results.append({'path': path.relative_to(ROOT).as_posix(), 'sha256': digest(path),
                        'pages': len(book), 'tables': len(document.tables), 'page_details': pages,
                        'contact_sheets': sheets, 'text_outside_page': out_of_page})
    for key, entry in provenance['outputs'].items():
        entry['sha256'] = digest(ROOT / entry['path'])
    provenance['rendering'] = {'engine': render['engine'], 'pdf_generated': False,
                               'page_counts': [result['pages'] for result in results]}
    provenance_path.write_text(json.dumps(provenance, ensure_ascii=False, indent=2) + '\n',
                                encoding='utf-8', newline='\n')
    record = {'content_and_input_checks': 'passed', 'docx_zip_integrity': 'passed',
              'template_placeholders': 'none', 'pdf_generated': False, 'documents': results,
              'visual_review': 'pending', 'schema_validation': 'pending after Word field update'}
    result_path = WEEK / '文档编制记录/验证结果.json'
    result_path.write_text(json.dumps(record, ensure_ascii=False, indent=2) + '\n',
                           encoding='utf-8', newline='\n')
    print(json.dumps({'documents': [{'path': row['path'], 'pages': row['pages'],
                                     'tables': row['tables'], 'sheets': row['contact_sheets']}
                                    for row in results]}, ensure_ascii=False))


if __name__ == '__main__':
    main()
