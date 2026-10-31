"""Build course-template DOCX review drafts from confirmed facts and repository evidence.

Requires python-docx. Reads source as text/AST only; never starts Django or exports PDF.
"""

import ast
import hashlib
import io
import json
import subprocess
import tokenize
import warnings
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[3]
WEEK = ROOT / 'doc/week7-8'
FACTS_PATH = WEEK / 'content/week78_confirmed_facts.json'
FACTS = json.loads(FACTS_PATH.read_text(encoding='utf-8'))
MEMBERS = FACTS['members']
SOURCE_REF = subprocess.check_output(['git', 'rev-parse', 'HEAD'], cwd=ROOT, text=True).strip()
SOURCE_URL = 'https://github.com/Adminer-2415304346/class-3-group-6'
OUTPUTS = {
    'coding': WEEK / 'AI Agent工程实践-Python语言编码规范.docx',
    'quality': WEEK / 'AI Agent工程实践-开源软件质量分析报告.docx',
}
RECORD = WEEK / '文档编制记录'
NAVY = '17365D'
GRAY = '526071'
AMBER = 'FFF2CC'

REVIEW_CHECKS = [
    ('RC01', '版本与归属', '确认分支、App、基线/改进提交；保留原作者和补修者；先同步develop。'),
    ('RC02', '阅读与注释', '模块、关键函数及核心块有职责/理由/边界说明；逐文件统计明确新增、修订与既有。'),
    ('RC03', '人工分析', '风格、设计、技能正例、缺陷四类观察均有路径/符号/可复现依据，候选风险单列。'),
    ('RC04', '扫描可比性', '前后工具版本、规则、目标版本、cwd、文件和排除范围一致；保存源码/配置哈希。'),
    ('RC05', '改进与回归', '至少一项适当改进有验证；必要时保留修复前失败证据；不以语法通过代替行为测试。'),
    ('RC06', '数据与状态', '0与未测分开；告警、误报、解析错误分别记录；测试分母/隔离边界和遗留项明确。'),
    ('RC07', '跨App契约', '核对参数/返回类型、权限、缓存键及事务时机；覆盖普通写入、批量与完整级联路径。'),
    ('RC08', '安全与副作用', '检查HTTP方法/身份/CSRF、输入归属、HTML上下文、密钥与随机值、外呼超时。'),
    ('RC09', '提交与材料', '只提交本轮范围及必要证据；无虚拟环境/密钥；Word、MD、代码版本及数字可追溯。'),
    ('RC10', '本人确认', '复核人填写日期、问题、处理及结论；未确认留待补；集成工具核查不代替本人签署。'),
]

PLAN_ROWS = [
    ('第7周任务阶段', '精读/三层注释、人工四类分析、基线与规范草稿', '各App主责；SZX统筹', '源码说明、四份素材、可复核baseline', '具体日历日期待确认'),
    ('第8周任务阶段', '适当改进、回归验证、同配置复扫及互查', '各App主责与互查搭档', '修复及测试证据、after、本人互查结论', '具体日历日期待确认'),
    ('当前收尾阶段一', 'SY补交；YFY补统一基线；三对互查记录闭环', 'SY、YFY及各复核人', '材料齐备、记录一致；新增修复经验证', '截止时间待确认'),
    ('当前收尾阶段二', '规范/质量Word定稿、优化计划确认、组长核对', 'SZX汇总；全员确认本人数据', '两份Word与最终代码/证据对应', '截止时间待确认'),
]

OPTIMIZATIONS = [
    ('P1', '收尾完整交付', 'SY补齐；YFY同配置基线；互查、日期和两份Word确认', '各成员；SZX统筹', '六人素材齐，原始证据可比，结论由本人确认'),
    ('P1', '文章与搜索失效', '统一文章批量发布/撤稿、普通删除与搜索索引的生命周期', 'SZX/GCY，协调WZY', '撤稿不再公开展示/搜索命中；相关缓存与索引验证'),
    ('P1', 'OAuth账号边界', 'state、隐藏ID归属、签名时效/一次性及邮箱归属策略', 'GZX/YFY，协调GCY', '模拟伪造/过期/重复使用及未激活账号，不越权关联'),
    ('P1', '上传与外部依赖', '上传授权/大小/格式/路径；显式搜索初始化与TLS异常策略', 'SZX/GCY', '越权/伪装/超限拒绝；无搜索服务时普通读取可用'),
    ('P1', '找回密码限频', '在secrets与成功消费基础上增加限流、失败次数及配额', 'YFY/GCY', '有效期、次数、重复/过期与正常用户流程均有断言'),
    ('P2', '查询和数据一致性', '浏览量原子更新、分类环校验、评论树与计数查询优化', 'SZX/WZY', '并发计数不丢失；环拒绝；同数据对比查询数及结果'),
    ('P2', '构建与运维', '生产Vite产物、运维权限/参数/失败反馈、ChatGPT适配', 'SY/SZX，协调GCY', '未授权无系统副作用；失败不伪装成功；产物可用'),
]

REFERENCES = [
    ('S1', 'PEP 8：项目一致性、缩进、命名和导入等基础建议', 'https://peps.python.org/pep-0008/'),
    ('S2', 'Django 5.2：事务与提交后回调', 'https://docs.djangoproject.com/en/5.2/topics/db/transactions/'),
    ('S3', 'Python 3.12：secrets安全随机值', 'https://docs.python.org/3.12/library/secrets.html'),
]


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def git_files(app):
    data = subprocess.check_output([
        'git', 'ls-tree', '-r', '-z', '--name-only', SOURCE_REF, '--', f'src/DjangoBlog/{app}',
    ], cwd=ROOT)
    return [name for name in data.decode('utf-8').split('\0') if name.endswith('.py')]


def collect_inventory():
    result = {}
    parse_warnings = []
    for member in MEMBERS:
        files = git_files(member['app'])
        items = []
        for name in files:
            source = (ROOT / name).read_text(encoding='utf-8-sig')
            category = 'production'
            if 'migrations' in Path(name).parts:
                category = 'migration'
            elif Path(name).name.startswith('test') or 'tests' in Path(name).parts:
                category = 'test'
            with warnings.catch_warnings(record=True) as captured:
                warnings.simplefilter('always', SyntaxWarning)
                tree = ast.parse(source)
            parse_warnings += [f'{name}: {warning.message}' for warning in captured]
            comments = [token for token in tokenize.generate_tokens(io.StringIO(source).readline)
                        if token.type == tokenize.COMMENT]
            items.append({
                'path': name, 'relative': Path(name).relative_to('src/DjangoBlog').as_posix(),
                'category': category, 'physical_lines': len(source.splitlines()),
                'nonempty_lines': sum(bool(line.strip()) for line in source.splitlines()),
                'classes': sum(isinstance(node, ast.ClassDef) for node in ast.walk(tree)),
                'functions': sum(isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef))
                                 for node in ast.walk(tree)),
                'comment_lines': len(comments),
                'signed': {person['tag']: sum(person['tag'] in token.string for token in comments)
                           for person in MEMBERS},
                'sha256': digest(ROOT / name),
            })
        result[member['app']] = items
    return result, parse_warnings


INVENTORY, PARSE_WARNINGS = collect_inventory()
TEST_EVIDENCE = WEEK / '成员成果/SZX/GCY集成/2026-10-31/after/manifest.json'
TEST_RECORD = json.loads(TEST_EVIDENCE.read_text(encoding='utf-8'))
assert TEST_RECORD['test_count'] == 357 and TEST_RECORD['failure_count'] == 0
for name, expected in TEST_RECORD['source_sha256'].items():
    assert digest(ROOT / name) == expected, f'Verified test source has changed: {name}'


def run_font(run, size=10.5, bold=False, color='202A35', code=False):
    run.font.name = 'Consolas' if code else 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')


def para(document, text, *, small=False, bold=False, color='202A35'):
    paragraph = document.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.line_spacing = 1.15
    fmt.space_after = Pt(3 if '编码规范' in document.core_properties.title else 5)
    fmt.widow_control = True
    fmt.keep_with_next = False
    run_font(paragraph.add_run(str(text)), 9 if small else 10.5, bold, color)
    return paragraph


def heading(document, text, level=1, page=False):
    paragraph = document.add_heading(text, level)
    fmt = paragraph.paragraph_format
    fmt.keep_with_next = True
    fmt.page_break_before = page
    fmt.space_before = Pt(12 if level == 1 else 8)
    fmt.space_after = Pt(6)
    for run in paragraph.runs:
        run_font(run, 14 if level == 1 else 11.5, True, NAVY)
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '黑体')
    return paragraph


def callout(document, text):
    paragraph = para(document, text)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), AMBER)
    paragraph._p.get_or_add_pPr().append(shd)
    return paragraph


def table(document, headers, rows, weights=None, size=9):
    rows = list(rows)
    width = (document.sections[-1].page_width - document.sections[-1].left_margin
             - document.sections[-1].right_margin) / 360000
    weights = weights or [1] * len(headers)
    widths = [width * weight / sum(weights) for weight in weights]
    grid = document.add_table(rows=1, cols=len(headers))
    grid.alignment = WD_TABLE_ALIGNMENT.CENTER
    grid.autofit = False
    for column, cm in zip(grid.columns, widths):
        column.width = Cm(cm)
    for index, values in enumerate([headers, *rows]):
        row = grid.rows[0] if index == 0 else grid.add_row()
        prop = row._tr.get_or_add_trPr()
        prop.append(OxmlElement('w:cantSplit'))
        if index == 0:
            prop.append(OxmlElement('w:tblHeader'))
        for cell, value, cm in zip(row.cells, values, widths):
            cell.width = Cm(cm)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margin = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                node = OxmlElement('w:' + side)
                node.set(qn('w:w'), '70')
                node.set(qn('w:type'), 'dxa')
                margin.append(node)
            tcpr.append(margin)
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                node = OxmlElement('w:' + side)
                node.set(qn('w:val'), 'single')
                node.set(qn('w:sz'), '4')
                node.set(qn('w:color'), 'CCD4DE')
                borders.append(node)
            tcpr.append(borders)
            if index == 0 or '待补' in str(value) or '待确认' in str(value):
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:fill'), 'E9EFF6' if index == 0 else 'FFF8E3')
                tcpr.append(shd)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.line_spacing = 1.05
            paragraph.paragraph_format.space_after = Pt(2)
            # Keep short tables whole; avoid lone headers and lone final rows in longer tables.
            paragraph.paragraph_format.keep_with_next = (
                index < len(rows) if len(rows) <= 5
                else index < 2 or index == len(rows) - 1
            )
            run_font(paragraph.add_run(str(value)), size, index == 0)
    gap = document.add_paragraph()
    gap.paragraph_format.space_after = Pt(0)
    gap.paragraph_format.line_spacing = Pt(2)
    run_font(gap.add_run(''), 2)
    return grid


def code(document, text):
    lines = text.strip('\n').splitlines()
    for index, line in enumerate(lines):
        paragraph = document.add_paragraph()
        fmt = paragraph.paragraph_format
        fmt.left_indent = Cm(0.25)
        fmt.right_indent = Cm(0.15)
        fmt.line_spacing = Pt(12)
        fmt.space_after = Pt(0)
        fmt.keep_with_next = index < len(lines) - 1
        fmt.keep_together = True
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:fill'), 'F3F5F7')
        paragraph._p.get_or_add_pPr().append(shd)
        run_font(paragraph.add_run(line or ' '), 8.5, code=True)
    para(document, '', small=True).paragraph_format.space_after = Pt(0)


def field(paragraph, instruction, text='1'):
    node = OxmlElement('w:fldSimple')
    node.set(qn('w:instr'), instruction)
    run = OxmlElement('w:r')
    value = OxmlElement('w:t')
    value.text = text
    run.append(value)
    node.append(run)
    paragraph._p.append(node)


def hyperlink(paragraph, label, url):
    from docx.opc.constants import RELATIONSHIP_TYPE
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK,
                                               is_external=True))
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    color = OxmlElement('w:color')
    color.set(qn('w:val'), '195595')
    properties.append(color)
    run.append(properties)
    text = OxmlElement('w:t')
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def new_document(kind, title):
    template = WEEK / '模板' / ('编码规范.docx' if kind == 'coding' else '质量分析报告.docx')
    document = Document(template)
    body = document._element.body
    for element in list(body):
        if element.tag != qn('w:sectPr'):
            body.remove(element)
    for level in range(1, 4):
        name = f'Heading {level}'
        if name not in document.styles:
            style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            style.base_style = document.styles['Normal']
            outline = OxmlElement('w:outlineLvl')
            outline.set(qn('w:val'), str(level - 1))
            style._element.get_or_add_pPr().append(outline)
    for section in document.sections:
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.3)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.6)
        section.right_margin = Cm(2.4)
        section.header_distance = Cm(0.9)
        section.footer_distance = Cm(0.9)
        paragraph = section.header.paragraphs[0]
        paragraph.text = ''
        run_font(paragraph.add_run(f'{FACTS["group"]} · DjangoBlog · 第7—8周'), 8.5,
                 color=GRAY)
        paragraph = section.footer.paragraphs[0]
        paragraph.text = ''
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_font(paragraph.add_run('审阅稿  |  第 '), 8.5, color=GRAY)
        field(paragraph, 'PAGE')
        run_font(paragraph.add_run(' 页 / 共 '), 8.5, color=GRAY)
        field(paragraph, 'NUMPAGES')
        run_font(paragraph.add_run(' 页'), 8.5, color=GRAY)
    document.styles['Normal'].font.size = Pt(10.5)
    document.styles['Normal'].font.name = 'Times New Roman'
    document.styles['Normal']._element.get_or_add_rPr().get_or_add_rFonts().set(
        qn('w:eastAsia'), '宋体')
    props = document.core_properties
    props.title = title
    props.subject = '第七、八周软件工程方法学课程任务'
    props.author = FACTS['group']
    props.last_modified_by = '石兆翔（组长整理）'
    props.version = '0.1'
    props.content_status = '小组审阅稿；待补项明确标识'
    props.created = datetime(2026, 10, 31, tzinfo=timezone.utc)
    props.modified = datetime.now(timezone.utc)
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(8)
    run_font(paragraph.add_run(title), 22, True, NAVY)
    para(document, '第7—8周 · v0.1 小组审阅稿', bold=True, color=GRAY)
    para(document, f'组号：{FACTS["group"]}    编制日期：2026年10月31日')
    para(document, '小组成员：' + '、'.join(member['name'] for member in MEMBERS))
    para(document, f'代码快照：develop@{SOURCE_REF[:7]}    汇总人：石兆翔', small=True)
    return document


def normalize_properties(document):
    orders = {
        'pPr': ('pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr '
                'suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap '
                'overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd '
                'snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc '
                'textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle '
                'rPr sectPr pPrChange').split(),
        'tcPr': ('cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar '
                 'textDirection tcFitText vAlign hideMark headers cellIns cellDel '
                 'cellMerge tcPrChange').split(),
    }
    for name, order in orders.items():
        ranks = {qn('w:' + value): index for index, value in enumerate(order)}
        for properties in document._element.body.iter(qn('w:' + name)):
            for child in sorted(list(properties), key=lambda item: ranks.get(item.tag, len(order))):
                properties.append(child)


def save(document, path):
    normalize_properties(document)
    temporary = path.with_name(path.stem + '.building.docx')
    document.save(temporary)
    temporary.replace(path)


def coding_document():
    document = new_document('coding', 'Python语言编码规范')
    callout(document, '本稿依据课程模板、现有工具配置与已确认案例整理。适用于本轮新增和修改代码；'
            '存量偏离逐项记录整改依据。本稿供全组审阅，成员本人互查结论另行确认。')
    heading(document, '1  规范制定的主要目的')
    para(document, '统一六人协作时的命名、格式、注释、接口及验证方式，使代码能被另一位成员理解、'
         '复现和维护；将质量问题落实到具体路径、触发、修复及证据。')
    para(document, '规范同时约束行为兼容和证据真实性。减少告警不能以引入功能回归为代价；'
         '缓存、权限和输入输出的契约必须由相应回归保护。')
    heading(document, '2  规范的适用范围')
    table(document, ['范围', '适用说明'], [
        ('六个主责App', 'blog、accounts、comments、oauth、servermanager、djangoblog；模板/前端按关联接口核查。'),
        ('代码类别', '业务Python、包入口、管理命令与模板标签；测试和迁移分别统计与审查。'),
        ('实施边界', '新增/修改代码优先满足；公共名称与返回类型改动先协调，存量整改不随意扩大补丁。'),
        ('优先依据', '课程要求、已确认业务契约、本组规则和统一工具配置；基础风格参考S1。'),
    ], [1, 4])
    heading(document, '3  规范')
    rules = [
        ('CS01  编码、缩进与排版',
         '使用UTF-8、4空格缩进，文件末尾保留换行，去除无意义行尾空格；本组行长为100字符，以ruff.toml为准。'
         '这是本组项目约定，不将PEP 8的默认建议误写为100。顶层定义之间留出合理空行，长表达式采用括号换行。'),
        ('CS02  命名与兼容',
         '模块、函数和变量使用有业务含义的snake_case，类使用CapWords，常量使用UPPER_CASE。'
         '不混用id、邮箱、第三方标识等不同概念；修改既有公开名称、路由、模型字段前核对调用与迁移，避免只为改名破坏兼容。'),
        ('CS03  导入与依赖',
         '按标准库、第三方、本地模块分组，生产代码使用显式导入，移除确实未使用的导入与变量。'
         '采用统一扫描cwd；未经核实不对全目录自动修复。避免模块导入时触发业务写库或网络外呼。'),
        ('CS04  函数职责与接口',
         '函数聚焦一个清晰职责，写明关键参数、返回值、状态及副作用。保留已约定返回类型，'
         '例如Article.comment_list仍返回QuerySet；不能为了查询优化擅自改成list，导致filter等调用失效。'),
        ('CS05  阅读注释与署名',
         '按模块层、关键函数/方法层、核心逻辑层说明职责、理由和边界，使用# SZX：等本人缩写；保留原作者，'
         '由他人修订时注明修订者。区分当前注释、既有注释、本周新增和修订，不能把总量当新增或替别人确认阅读。'),
        ('CS06  可选参数与异常',
         '先识别可选上下文，再取值或调用；不以bare except掩盖索引、类型或业务错误。只处理能够合理恢复的异常，'
         '必要时记录日志。可选get_cache_key缺失可回退，但方法内部真异常应向上传播。G2的无参回归作为反例保留证据。'),
        ('CS07  ORM、归属与状态',
         '对公开入口按发布、启用、激活等真实状态筛选；回复父ID、隐藏字段和第三方ID必须校验归属。'
         '校验失败返回约定的表单错误或HTTP状态，不让已知无效输入冒泡为500；模型choices不能冒充数据库强约束。'),
        ('CS08  写入权限与请求边界',
         '有副作用入口明确HTTP方法、身份/权限及CSRF，不能用GET执行清缓存等管理写操作。'
         '直接login与authenticate路径分别核对；不能根据前端隐藏字段或cookie声称服务端已完成权限判断。'),
        ('CS09  模板与HTML上下文',
         '区分普通文本、HTML、属性和脚本上下文，优先自动转义或适当框架接口。mark_safe不执行净化。'
         'JSON嵌入单引号HTML属性时仍须按属性上下文强制转义；昵称中的引号和尖括号必须能安全往返。'),
        ('CS10  缓存、事务与批量操作',
         '记录真实缓存键、空值语义、所有读写入口及失效时机。数据库相关展示缓存成功提交后失效，回滚保留[S2]；'
         '批量更新固定待改主键并按能力锁行，显式调用失效入口。完整级联删除按实际被删对象覆盖各文章，不能只取一层子回复。'
         '保留无关验证码等状态，不以cache.clear替代局部失效。'),
        ('CS11  外部服务与通知',
         'HTTP显式设置超时，记录失败和降级边界；本项目OAuth使用10秒、百度推送5秒的现有约定。'
         'TLS校验与错误处理不可默默降级；真实联调和mock结果分别说明。缓存清理不附带邮件，事务通知时机与投递保证分开约定。'),
        ('CS12  密钥与验证码',
         '密钥来自配置/环境，报告与日志不含真实凭据。找回密码等安全随机值使用secrets[S3]；当前六位ASCII数字协议'
         '保留前导零与重复数字。随机值、有效期、成功消费、限频和失败次数是不同控制点；未实现项列入后续计划。'),
        ('CS13  测试与查询预算',
         '验证正常、无效、匿名/越权、空值及关键事务/并发交错边界，断言明确结果。缺陷回归在可行时留修复前失败证据。'
         'R4测量边界为单次get_reactions_summary，当前1查询并满足最多2查询预算，不将整页成本混入；不得漏计人数或名单外用户状态。'),
        ('CS14  扫描与证据',
         '固定Ruff0.17.1、Bandit1.9.4、规则与目标版本、cwd和生产/测试/迁移范围。保存原始JSON、命令、退出码、源码/配置哈希。'
         '发现、误报、解析错误与未执行分开；Ruff返回1表示发现告警，不等于工具崩溃。旧证据不覆盖，新轮次另存。'),
        ('CS15  协作、复核与提交',
         '开工前fetch并核对个人分支是否已整合最新develop；拉取远端引用不等于当前分支完成合并。'
         '跨App先确认接口和范围，冲突保留双方有效逻辑及作者。个人分支提交到develop，全组完成后由组长负责main最终合并。'
         '本人复核日期、意见和结论由本人确认，自动化核查不代签。'),
    ]
    for title, text in rules:
        heading(document, '3.' + title[2:4].lstrip('0') + '  ' + title[6:], 2)
        para(document, f'[{title[:4]}] ' + text)
        if title.startswith('CS06'):
            code(document, '''key = None
if args:
    get_cache_key = getattr(args[0], 'get_cache_key', None)
    if callable(get_cache_key):
        key = get_cache_key()  # 方法自身错误传播''')
            para(document, '摘录：djangoblog/utils.py:cache_decorator的可选上下文处理；'
                 '完整回退实现及回归见仓库对应方法。', small=True)
        if title.startswith('CS09'):
            code(document, "data-reactions='{{ comment_item|get_reactions_for_user:user|to_json|force_escape }}'")
            para(document, '实例：templates/comments/tags/comment_item_modern.html；实际SSR属性往返测试覆盖恶意昵称。', small=True)
        if title.startswith('CS12'):
            code(document, "return ''.join(secrets.choice(string.digits) for _ in range(6))")
            para(document, '实例：djangoblog/utils.py:generate_code；本轮账户回归保留成功消费验证码的验证。', small=True)
    heading(document, '3.16  本轮互查通过条件', 2)
    table(document, ['编号', '检查项', '通过条件'], REVIEW_CHECKS, [0.7, 1.1, 4], 8.8)
    para(document, 'P1：功能/权限回归、证据失真或范围错误，修复并验证后再验收；P2：影响可靠性或维护的缺项，'
         '明确整改/保留理由；P3：低影响排版与历史风格项，纳入有证据的后续计划。通过可附遗留项，但不能隐去阻断问题。')
    heading(document, '3.17  本组可执行工具约定与参考', 2)
    code(document, '''python doc/week7-8/tools/scan_app.py <app> baseline
python doc/week7-8/tools/scan_app.py <app> after --output-dir <新证据目录>''')
    para(document, '静态工具使用独立环境；cwd由统一驱动固定为src/DjangoBlog；业务扫描排除migrations、test*.py、tests目录。'
         '测试/迁移仍人工阅读并单列统计，不从报告中消失。')
    for label, name, url in REFERENCES:
        paragraph = para(document, f'[{label}] ', small=True)
        hyperlink(paragraph, name, url)
    para(document, '项目依据：tools/ruff.toml、tools/bandit.yaml、个人任务单、已确认源码及成员成果；'
         '外部资料访问日期为2026-10-31。具体代码证据以质量报告附录索引为准。', small=True)
    save(document, OUTPUTS['coding'])


def production_items(app):
    return [item for item in INVENTORY[app] if item['category'] == 'production']


def quality_document():
    document = new_document('quality', '“DjangoBlog”开源软件质量分析报告')
    callout(document, '阶段稿：已填入5人成果、357项已通过隔离回归及已核对数据；SY本轮素材、YFY统一基线、'
            '互查确认和双周日期仍待补。本文不是全组最终验收完成声明。')
    heading(document, '1  待分析开源软件概述')
    para(document, 'DjangoBlog是以Python/Django实现的技术博客系统，提供文章与分类标签、站点设置、账号认证、第三方授权、'
         '评论/回复与表情反应、检索，以及运维和邮件日志等功能。关联展示采用Django模板和前端资源。')
    para(document, f'本轮主评范围是六个已注册的自有App。统计从Git提交{SOURCE_REF}读取跟踪文件，按AST和token计数；'
         '依赖库、虚拟环境、前端文件与生成文档不混入生产Python分母。类包括嵌套类，函数数含方法；物理行不等于有效代码行。')
    rows = []
    for member in MEMBERS:
        items = production_items(member['app'])
        rows.append((member['app'], len(items), sum(item['physical_lines'] for item in items),
                     sum(item['classes'] for item in items), sum(item['functions'] for item in items)))
    rows.append(('合计', *[sum(row[column] for row in rows) for column in range(1, 5)]))
    table(document, ['App', '生产Python文件', '物理行', 'AST类', '函数/方法'], rows, [1.6, 1.2, 1, 0.8, 1])
    table(document, ['App', '测试/测试基类文件', '迁移Python文件'], [
        (member['app'], sum(item['category'] == 'test' for item in INVENTORY[member['app']]),
         sum(item['category'] == 'migration' for item in INVENTORY[member['app']]))
        for member in MEMBERS
    ], [1.5, 2, 2])
    para(document, 'SY模块的仓库规模可以静态核实，但当前未见其本轮个人成果，不能据规模表认定其完成阅读、注释或改进。')
    heading(document, '2  软件质量分析方法及工具')
    heading(document, '2.1  分析的软件质量要素', 2)
    table(document, ['要素', '本轮关注点'], [
        ('内部质量', '风格与命名、注释真实性、模块职责、接口兼容、异常与缓存设计、查询成本。'),
        ('外部质量', '发布/启用/激活边界、权限与输入归属、异常响应、HTML安全、批量与级联一致性。'),
        ('可追溯性', '固定基线/配置/分母与源码指纹；测试和原始结果可定位，真实联调和隔离测试分开。'),
    ], [1, 4])
    heading(document, '2.2  软件质量分析方法', 2)
    para(document, '按App阅读入口、核心对象与调用链，在模块/方法/关键块署名说明；从风格、设计、技能正例与问题四类人工分析。'
         '静态发现结合源码核实，选取至少一项适当改进，保留必要的失败复现、修复、回归和同配置复扫。')
    para(document, '互查按RC01—RC10执行。候选风险不直接写成既成缺陷；集成验证与成员本人确认分别记录。'
         'SY未提交、YFY配置变更前后不可直接对比及双方互查空白均保留为待补。')
    heading(document, '2.3  软件质量分析工具', 2)
    table(document, ['工具', '用途与边界'], [
        ('Ruff 0.17.1', 'E/F/I/B/C90，行长100，复杂度10，统一target=py312；静态告警不证明业务结果。'),
        ('Bandit 1.9.4', '保留全部规则并ignore-nosec；生产显式文件清单排除测试/迁移，误报结合上下文分流。'),
        ('Django 5.2.17测试', '使用SQLite内存、LocMem缓存/邮箱和SimpleEngine；阻断HTTP/socket，覆盖选定标签。'),
        ('Git、AST、tokenize', '核对提交/合并、逐文件规模、结构和注释token；哈希记录源码和配置，计数不是作者确权。'),
    ], [1.4, 4])
    heading(document, '3  人工分析代码的质量情况')
    heading(document, '3.1  代码风格', 2)
    para(document, 'G1将admin_site星号导入改为显式来源；G3清理Whoosh未用导入和变量；WZY最终相关业务11文件Ruff/Bandit均0。'
         '同时，blog和派生检索实现仍有历史排版、导入及复杂度项，不能把部分范围零告警写成整个系统符合规范。')
    para(document, '判定依据采用《Python语言编码规范》CS01—CS05与统一配置。按原作者及SZX补修分别记贡献，'
         '本周新增统计不能由当前缩写行数倒推；逐文件当前快照见附录B。')
    heading(document, '3.2  代码设计', 2)
    table(document, ['观察', '源码与结论'], [
        ('职责与接口', '六App按业务划分；Article.comment_list的QuerySet返回与消费者filter接口保持。'),
        ('缓存调用契约', 'cache_decorator兼容无参函数/参数调用/自定义键；方法真实错误传播，避免G2收窄异常引入IndexError。'),
        ('完整生命周期', '评论pre_save记录旧文章，post_save/post_delete在实际写库连接on_commit失效；批量动作显式失效。'),
        ('数据与整页边界', '归档保留文章数据缓存，HTML按请求渲染，侧栏不被旧整页响应遮挡；外部CDN/浏览器旧响应不在本轮保证内。'),
        ('通知与异常隔离', '缓存入口不发邮件；提交后独立通知，回滚不发送；无__name__插件回调出错仍继续后续链。'),
    ], [1.1, 4])
    heading(document, '3.3  编程技能', 2)
    para(document, '已核对实例包括显式接口与状态筛选、select_related等预加载、回调/Filter扩展、白名单HTML处理、'
         '事务提交后失效、必要字段一次读取聚合与原始证据指纹。每个正例均有适用边界，不能只引用mark_safe、choices或缓存命中即声称安全。')
    code(document, '''# djangoblog/utils.py：可选接口与真实异常分开
key = None
if args:
    get_cache_key = getattr(args[0], 'get_cache_key', None)
    if callable(get_cache_key):
        key = get_cache_key()''')
    para(document, 'R4同一个SELECT计算count、最多10名显示用户和当前查看者参与态。八种已参与表情场景中，'
         '历史登录17/匿名9次，WZY初版2/2，本轮补修1/1；0/1/8种、名单外用户、SSR/API与交错增删均有回归。'
         '仍需遍历该评论全部反应；数据库驱动可能缓冲结果，不宣称无限规模常数内存。')
    heading(document, '3.4  代码质量问题', 2)
    table(document, ['问题', '处理与当前结论', '归属/证据'], [
        ('草稿公开、未授权清缓存、HTML输出', '已按发布/权限/上下文修复并回归；源阶段分析见README。', 'SZX；EV01'),
        ('空用户名/未激活认证与验证码重放', '后端加固、重置成功后消费验证码；84项账户相关测试通过。', 'YFY；EV02'),
        ('OAuth外呼无超时', '四处显式timeout=10；26项业务测试通过。state/IDOR/邮箱归属等仍在后续计划。', 'GZX；EV03'),
        ('跨文章/无效父回复及关闭评论异常', 'R1—R3视图校验修复，状态与表单反馈回归保留。', 'WZY；EV04'),
        ('R4两次读取与R5漏清后代缓存', 'SZX补为单次读取与完整生命周期；318项阶段回归通过，357项集成保持。', 'WZY原改动+SZX补修；EV05'),
        ('G2无参调用IndexError', '原22项全过，GCY初版3项报错；SZX显式识别可选上下文后关闭回归。', 'GCY原改动+SZX补修；EV06'),
        ('验证码用途/插件错误日志', 'SZX改secrets保持六位协议；回调无__name__不阻断后续，新增最小回归通过。', 'SZX；EV06'),
        ('SY运维模块', '待补：本轮人工分析、实际改进、验证、扫描及互查，不能按既有源码认定完成。', 'SY待提交'),
    ], [1.4, 3.1, 1.4], 8.8)
    heading(document, '4  工具分析代码的质量情况')
    para(document, '下表保留各阶段已确认结果及其分母。不同配置、范围或代码时点的结果不合计为“全系统下降”。'
         '告警数量不是缺陷数量；没有执行不能填0。当前集成源码与357项通过记录逐文件哈希对应。')
    table(document, ['App/阶段', '生产文件', 'Ruff', 'Bandit', '可比性与来源'], [
        ('blog原阶段', '22', '66→65', '45→39', '同阶段可比；当前新增缓存入口后生产23，历史扫描不能冒充完整最新范围。EV01'),
        ('accounts', '10', '10→8*', '0→0*', '*基线target313，最终target312；同配置baseline待补，不做净下降结论。EV02'),
        ('oauth', '10', '12→12', '10→6', 'B113×4关闭；B105 URL/空串误报另记。前后配置一致。EV03'),
        ('comments最终', '11', '0', '0', '最终补修快照；旧基线9及中间3/4保留，SZX执行者单列。EV04/05'),
        ('djangoblog统一阶段', '22', '72→34→28', '8→7→5', '原版→GCY→SZX集成，包含上游修复；原自建混入测试数字另存。EV06'),
        ('servermanager', '11现状', '待交', '待交', '这里只确认静态规模，未取得本轮个人工具证据。'),
    ], [1.3, 0.7, 0.9, 0.9, 3], 8.5)
    para(document, 'djangoblog原自建Bandit实际扫27文件，含5个测试文件，after的38条中31条来自测试，'
         '不能按22个生产文件口径合计；已转统一入口并另存after-szx-integration。当前剩余Ruff I001×7、E501×14、'
         'C901×5、B007×1、B904×1；Bandit为日志链接B308/B703各2及Whoosh运行时assert B101×1，处置理由保留。')
    heading(document, '4.1  已确认行为验证与边界', 2)
    table(document, ['范围', '数量', '结果'], [
        ('blog所选标签', '158', '通过'), ('accounts所选标签', '84', '通过'),
        ('comments所选标签', '50', '通过'), ('oauth业务逻辑标签', '26', '通过'),
        ('djangoblog既有22项+新增17项', '39', '通过'), ('合计', '357', 'failure_count=0'),
    ], [3, 0.8, 1.2])
    para(document, '环境：Python3.12.7、Django5.2.17、SQLite内存、LocMem缓存/邮箱、SimpleEngine，HTTP/socket阻断。'
         '不代表全仓库全部测试、真实MySQL/Redis/CDN、OAuth提供商、SMTP投递或前端生产构建均已验证。'
         '本次文档编制核对源码与已有357项通过记录一致，没有因整理Word虚构重新执行测试。')
    heading(document, '5  小组分工')
    table(document, ['成员/标识', '分析模块', '主要工作与注释范围', '互查/交付状态'], [
        (f'{m["name"]}\n{m["id"]}\n# {m["tag"]}：', m['app'],
         f'生产Python {len(production_items(m["app"]))}文件；当前逐文件注释见附录B。' +
         ('本轮个人阅读/新增统计、改进证据待交。' if m['tag'] == 'SY' else '人工分析、工具和验证按本人材料及SZX补修分别记录。'),
         m['status'] + f'；搭档{m["peer"]}') for m in MEMBERS
    ], [1.3, 0.9, 2.5, 2], 8.8)
    para(document, 'SZX的组长/集成补修不计作其他成员本人返工或人工签署。WZY/GCY原提交在个人分支保留，'
         '后续修复从szx_branch进入develop；不能根据个人分支旧版本再次认定集成版本仍有已修复回归。')
    heading(document, '6  下一步的改进计划')
    table(document, ['级别', '方向', '拟实施内容', '建议负责人', '验收条件'], OPTIMIZATIONS,
          [0.5, 1, 2.3, 1.1, 2.1], 8.5)
    para(document, '以上后续优化计划依据README和已确认剩余项整理；优先级与责任为拟定收尾/改进安排，'
         '最终截止时间待组长确认，不能写成全部已实施或每位成员已承诺完成。')
    heading(document, '附录A  组会、双周计划与互查标准', page=True)
    meeting = FACTS['meeting']
    table(document, ['项目', '已确认内容'], [
        ('组会时间', f'{meeting["date"]} {meeting["start"]}—{meeting["end"]}'),
        ('形式/地点', f'{meeting["mode"]}，{meeting["place"]}'),
        ('出席/缺席', f'{meeting["attendance"]}；缺席：{meeting["absent"]}'),
        ('主持/记录', f'{meeting["chair"]} / {meeting["recorder"]}'),
        ('记录依据', '组长确认会议事实及结论沿用既定沟通安排；本节为补录，不是逐字记录。'),
    ], [1, 4])
    for index, conclusion in enumerate(meeting['conclusions'], 1):
        para(document, f'结论{index}：{conclusion}')
    para(document, '10月31日的代码集成、测试和补修属于后续进度快照，见正文，不倒填为10月23日当时完成的事项。', small=True)
    heading(document, 'A.1  双周任务阶段与当前收尾安排', 2)
    table(document, ['阶段', '任务', '责任', '交付/验收', '时间'], PLAN_ROWS,
          [1, 2.2, 1.3, 2.1, 1], 8.5)
    plan = FACTS['plan']
    para(document, f'双周起始：{plan["start_date"] or "【待确认】"}；结束：{plan["end_date"] or "【待确认】"}；'
         f'最终截止：{plan["deadline"] or "【待确认】"}。不以本次编制日期倒推实际计划日期。')
    heading(document, 'A.2  互查标准与闭环状态', 2)
    table(document, ['编号', '检查项', '通过条件'], REVIEW_CHECKS, [0.7, 1.1, 4], 8.8)
    table(document, ['配对', '现有事实', '待确认/待补'], [
        ('SZX↔GCY', 'SZX侧集成工具核查和修复证据已入库。', 'GCY对SZX的本人反向互查。'),
        ('YFY↔GZX', '双方各有单侧核查内容，YFY第8节已有对GZX评审。', 'GZX互查表和YFY第7节仍有空白，需同步处理及本人确认。'),
        ('WZY↔SY', 'WZY材料与SZX补修已集成。', 'SY提交本轮成果；双方本人复核、处理和结论。'),
    ], [1, 2.2, 2.8], 8.8)
    heading(document, '附录B  当前集成源码的逐文件注释清单', page=True)
    para(document, '口径：业务Python中的COMMENT token物理行；含缩写按注释字符串包含该缩写计数，可能同一行含多个标识，'
         '不作为独占署名或本周新增量。测试/迁移不混入；本周新增/修订需由个人差分证据确认。此清单填入可核实的当前事实。')
    for member in MEMBERS:
        heading(document, f'B.{MEMBERS.index(member) + 1}  {member["name"]} / {member["app"]}', 2)
        if member['tag'] == 'SY':
            callout(document, '待补：SY尚未交本轮逐文件阅读/新增注释记录；不据现有代码补写其个人完成情况。'
                    f'现状生产文件{len(production_items(member["app"]))}个已计入概述规模表。')
            continue
        rows = []
        for item in production_items(member['app']):
            if item['comment_lines'] or item['signed'][member['tag']]:
                others = '、'.join(f'{tag}:{count}' for tag, count in item['signed'].items()
                                  if tag != member['tag'] and count)
                rows.append((item['relative'], item['comment_lines'], item['signed'][member['tag']],
                             others or '—'))
        table(document, ['文件（相对src/DjangoBlog）', '注释token行', f'含{member["tag"]}', '其他缩写'],
              rows, [3, 1, 0.8, 1.4], 8.5)
        para(document, '未列出的该App业务文件没有COMMENT token；当前总量不等于本周新增，'
             f'本人周际统计与来源见成员成果/{member["tag"]}/注释统计或SZX阶段验证。', small=True)
    heading(document, '附录C  证据索引及定稿待补项', page=True)
    table(document, ['编号', '仓库路径/来源', '用途'], [
        ('EV01', 'doc/week7-8/README.md；成员成果/SZX/原始工具结果/quality-summary.json', 'blog原阶段分析、规模和工具对比'),
        ('EV02', '成员成果/YFY/工具与验证记录.md；原始工具结果/；验证日志/', 'accounts材料、84项；同配置baseline待补'),
        ('EV03', '成员成果/GZX/工具与验证记录.md、互查记录.md及原始工具结果/', 'oauth超时、26项和互查状态'),
        ('EV04', '成员成果/WZY/四份素材及after-r4-r5-szx/', 'comments分阶段问题与最终11文件复扫'),
        ('EV05', '成员成果/SZX/集成修复/2026-10-31/验证结果.md', 'R4/R5边界修复与318项阶段验证'),
        ('EV06', '成员成果/SZX/GCY集成/2026-10-31/验证结果.md；after/manifest.json', 'GCY补修、357项与逐文件指纹'),
        ('EV07', '成员成果/GCY/四份素材及after-szx-integration/', 'djangoblog22文件统一复扫与修订统计'),
        ('EV08', 'content/week78_confirmed_facts.json；第7—8周组会、双周计划与互查标准.md', '组长确认事实、计划待补日期及当前清单'),
        ('EV09', f'GitHub PR37；develop@{SOURCE_REF[:7]}', '5人成果已集成，保留GCY原作者历史'),
    ], [0.6, 3.8, 1.6], 8.3)
    para(document, '除EV01首项README外，上表“成员成果/”路径均相对doc/week7-8。完整源码/模板/输入指纹见文档编制记录。', small=True)
    table(document, ['定稿项', '当前状态', '责任/确认方'], [
        ('SY个人交付', '待补：源码、四份素材、改进/验证/复扫、互查。', 'SY；WZY复核'),
        ('YFY同配置baseline', '待补；保留已披露配置变化，不用旧10→8作统一下降。', 'YFY'),
        ('本人互查闭环', '待补空白、不同步状态及本人结论；不代签。', '三对搭档'),
        ('双周具体日期', '待确认起止与最终截止；任务阶段和责任已列。', '组长'),
        ('Word全组审阅', '本稿已填已确认内容；收齐数据后统一版本和确认状态。', 'SZX汇总；全员'),
    ], [1.1, 3.1, 1.3], 9)
    save(document, OUTPUTS['quality'])


def supporting_markdown():
    meeting = FACTS['meeting']
    text = f'''# 第7—8周组会、双周计划与互查标准

编制：2026-10-31；{FACTS['group']}；组长石兆翔整理。代码进度快照develop@{SOURCE_REF[:7]}。
会议事实及既定安排由组长在对话中确认；细化的验收清单与当前收尾安排在本次整理，不代填成员本人复核。

## 1. 已确认组会记录

- 时间：{meeting['date']} {meeting['start']}—{meeting['end']}。
- 形式/地点：{meeting['mode']}，{meeting['place']}。
- 出席：{meeting['attendance']}；缺席：{meeting['absent']}。
- 主持/记录：{meeting['chair']} / {meeting['recorder']}。
- 记录说明：据组长确认补录，不是逐字记录；后续10月31日修复结果另列为进度快照。

### 会议结论（按组长确认的沟通安排整理）

'''
    text += '\n'.join(f'{index}. {value}' for index, value in enumerate(meeting['conclusions'], 1))
    text += '\n\n## 2. 六人职责与互查配对\n\n| 成员 | App | 个人分支 | 互查搭档 |\n| --- | --- | --- | --- |\n'
    text += '\n'.join(f'| {m["name"]} {m["tag"]} | {m["app"]} | {m["branch"]} | {m["peer"]} |' for m in MEMBERS)
    text += '\n\n## 3. 双周阶段与当前收尾计划\n\n| 阶段 | 任务 | 责任 | 交付/验收 | 时间 |\n| --- | --- | --- | --- | --- |\n'
    text += '\n'.join('| ' + ' | '.join(row) + ' |' for row in PLAN_ROWS)
    text += '\n\n双周起止与最终截止日期：**待组长确认**；不推算或回填虚构日期。\n'
    text += '\n## 4. 统一互查标准\n\n| 编号 | 检查项 | 通过条件 |\n| --- | --- | --- |\n'
    text += '\n'.join('| ' + ' | '.join(row) + ' |' for row in REVIEW_CHECKS)
    text += '''

每项记录：基线/改进提交、文件与符号、证据、问题级别、本人处理、复测、复核日期与结论。
P1阻断项修复并验证后再验收；P2/P3明确整改或保留理由。工具核查与本人确认分别记录。

## 5. 当前闭环状态

- SZX↔GCY：SZX集成核查完成；GCY对SZX的本人反向互查待确认。
- YFY↔GZX：双方有单侧核查内容，仍需同步GZX互查表与YFY第7节空白，并由本人确认。
- WZY↔SY：WZY集成补修已验证；SY提交本轮材料后完成双方本人复核。
- 不把留白改为“已通过”，不把个人分支旧代码误认为当前develop仍未修复。

## 6. 当前进度与后续交付

5人成果已合入develop；357项选定范围隔离测试通过并核对源码指纹。SY本轮成果、YFY统一基线、互查确认、双周日期与两份Word最终确认仍待收齐。
编码规范已形成审阅稿，质量报告框架已填入确认数据。下一步优化按优先级/责任/验收条件列入报告第6节。
个人分支提交到develop；全组完成后由组长负责main最终合并。本文不代签任何成员本人确认。
'''
    path = WEEK / '第7—8周组会、双周计划与互查标准.md'
    path.write_text(text.rstrip() + '\n', encoding='utf-8', newline='\n')


def main():
    RECORD.mkdir(exist_ok=True)
    coding_document()
    supporting_markdown()
    quality_document()
    input_paths = [FACTS_PATH, Path(__file__), WEEK / '模板/编码规范.docx', WEEK / '模板/质量分析报告.docx',
                   TEST_EVIDENCE, WEEK / 'tools/ruff.toml', WEEK / 'tools/bandit.yaml']
    manifest = {
        'compiled_date': FACTS['compiled_date'], 'source_ref': SOURCE_REF,
        'purpose': 'Course Word review drafts with confirmed content; not final group acceptance.',
        'meeting_confirmed': FACTS['meeting'], 'plan_dates': FACTS['plan'],
        'test_evidence': {'test_count': 357, 'failure_count': 0,
                          'verified_code_commit': TEST_RECORD['verified_code_commit']},
        'inputs_sha256': {path.relative_to(ROOT).as_posix(): digest(path) for path in input_paths},
        'inventory': INVENTORY, 'source_parse_warnings': PARSE_WARNINGS,
        'outputs': {kind: {'path': path.relative_to(ROOT).as_posix(), 'sha256': digest(path)}
                    for kind, path in OUTPUTS.items()},
        'rendering': 'Pending native Word XPS-to-PNG verification; no PDF deliverable.',
    }
    (RECORD / '编制与来源.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n',
                                          encoding='utf-8', newline='\n')
    print(json.dumps({'source_ref': SOURCE_REF,
                      'outputs': {key: str(path) for key, path in OUTPUTS.items()}}, ensure_ascii=False))


if __name__ == '__main__':
    main()
