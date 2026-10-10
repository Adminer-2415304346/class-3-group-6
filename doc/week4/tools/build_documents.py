"""Fill the two provided course DOCX templates with verified project content."""

from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
OUT = Path(__file__).resolve().parents[1]
ASSETS = OUT / 'assets'
TEMPLATE_DIR = ROOT.parent
UI_TEMPLATE = TEMPLATE_DIR / '文档模板-软件界面设计说明书模板.docx'
DATA_TEMPLATE = TEMPLATE_DIR / '文档模板-软件数据模型设计说明书模板.docx'
BLUE = RGBColor(23, 49, 91)
GRAY = RGBColor(82, 99, 119)


def east_asia_font(style, name='Microsoft YaHei'):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    rfonts.set(qn('w:eastAsia'), name)


def set_cell_fill(cell, color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), color)
    tcpr.append(shd)


def keep_row_together(row):
    trpr = row._tr.get_or_add_trPr()
    item = OxmlElement('w:cantSplit')
    trpr.append(item)


def add_page_number(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' PAGE '
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t')
    text.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)


def prepare(template, title, kind):
    if not template.exists():
        raise FileNotFoundError(f'Course template missing: {template}')
    doc = Document(template)
    body = doc._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)

    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(.76)
    section.bottom_margin = Inches(.72)
    section.left_margin = Inches(.77)
    section.right_margin = Inches(.77)
    section.header_distance = Inches(.35)
    section.footer_distance = Inches(.4)

    normal = doc.styles['Normal']
    east_asia_font(normal)
    normal.font.size = Pt(10)
    normal.font.color.rgb = BLUE
    normal.paragraph_format.line_spacing = 1.32
    normal.paragraph_format.space_after = Pt(5)
    for name, size, before, after in [('Heading 1', 20, 0, 15), ('Heading 2', 15, 17, 8), ('Heading 3', 11.5, 11, 5)]:
        style = doc.styles[name]
        east_asia_font(style)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    doc.core_properties.title = title
    doc.core_properties.subject = kind
    doc.core_properties.keywords = '软件工程方法学,第四周,AI Agent 工程实践'
    header = section.header.paragraphs[0]
    header.text = '软件工程方法学  ·  第四周实践任务'
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = GRAY
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ')
    add_page_number(footer)
    footer.add_run(' 页')
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = GRAY

    doc.add_paragraph(title, 'Heading 1')
    p = doc.add_paragraph('24 级软件工程 3 班第 6 组  |  编制日期：2026-10-08  |  版本：个人分支初稿')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph('学号、参会人员及最终审核信息：待小组填写；本稿依据当前代码与本地演示环境编写。')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.color.rgb = GRAY
    return doc


def para(doc, text):
    return doc.add_paragraph(text)


def h2(doc, text):
    return doc.add_paragraph(text, 'Heading 2')


def h3(doc, text):
    return doc.add_paragraph(text, 'Heading 3')


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, value in enumerate(headers):
        t.rows[0].cells[i].text = value
        set_cell_fill(t.rows[0].cells[i], 'E7EFFB')
    header_flag = OxmlElement('w:tblHeader')
    t.rows[0]._tr.get_or_add_trPr().append(header_flag)
    keep_row_together(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        keep_row_together(t.rows[-1])
    for row_index, row in enumerate(t.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    run.font.name = 'Microsoft YaHei'
                    run.font.size = Pt(8.3)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = BLUE
    return t


def figure(doc, filename, caption, width=6.55):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(ASSETS / filename), width=Inches(width))
    p = doc.add_paragraph(caption)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(13)
    for run in p.runs:
        run.font.size = Pt(8.5)
        run.font.color.rgb = GRAY


def build_ui():
    doc = prepare(UI_TEMPLATE, 'AI Agent 工程实践博客系统的界面设计', '软件界面设计说明书')
    h2(doc, '博客系统介绍')
    para(doc, '本系统是在开源 DjangoBlog 基础上二次开发的内容发布与交流平台。第四周以“AI Agent 工程实践”为博客内容主题，围绕架构、工具与工作流、知识与记忆、评测与安全、应用案例组织文章。系统本身仍是博客，不声称提供 Agent 自动执行、文档上传或任务编排能力。')
    para(doc, '界面目标是让读者迅速发现主题文章、按栏目与标签浏览、阅读正文并参与评论；让编辑通过 Django Admin 维护内容。桌面与移动端共享内容结构，提供搜索入口、浅色/深色切换和可访问的导航反馈。')

    h2(doc, '功能介绍')
    table(doc, ('角色 / 场景', '主要界面与路径', '现有功能与关键元素'), [
        ('访客发现内容', '首页 /；分类 /category/<slug>.html；标签 /tag/<slug>.html', '栏目导航、文章摘要、发布时间、标签、侧栏推荐、分页'),
        ('访客查找内容', '搜索 /search/?q=关键词；归档 /archives.html', '关键词输入、结果高亮、按时间查看文章'),
        ('访客阅读', '文章 /article/年/月/日/id.html', '标题、面包屑、目录、Markdown 正文、上一篇/下一篇、评论列表'),
        ('注册用户互动', '文章评论区；登录 /login/；注册 /register/', '登录后发表评论、回复和表情反应；未登录时显示登录引导'),
        ('编辑管理', '管理后台 /admin/', '文章、分类、标签、站点配置等管理；需相应权限'),
    ])
    para(doc, '主题测试数据包含 12 篇文章、5 个子分类、14 个标签、3 条可见示例评论；这些数据只用于本地演示，不是用户真实投稿。')

    h2(doc, '主要的界面设计')
    h3(doc, '信息架构与界面流转')
    figure(doc, '06-ui-mindmap.png', '图 1  界面思维导图：主题栏目、阅读、互动与管理入口')
    figure(doc, '07-ui-flow.png', '图 2  主要界面流转：发现内容—阅读—登录—互动')
    para(doc, '首页可以直接进入文章详情，也能先按分类、标签或搜索结果筛选。详情页的面包屑返回栏目；发表评论需登录。管理员入口不属于访客主路径。图中的页面关系是现有路由与模板的概括，搜索索引需先构建才会返回本地文章。')

    h3(doc, '实际运行界面与核心元素')
    figure(doc, '01-home-desktop.png', '图 3  桌面首页：站点主题、栏目导航、最近文章与侧栏')
    figure(doc, '03-category.png', '图 4  分类页：按“基础与架构”等栏目筛选文章')
    figure(doc, '04b-search.png', '图 5  搜索页：关键词 Agent 的结果与标题高亮')
    figure(doc, '02-article-detail.png', '图 6  文章详情：面包屑、标题、目录、正文与侧栏')
    figure(doc, '02b-article-comments.png', '图 7  文章评论区：已发表的主题评论及未登录引导')
    figure(doc, '04-login.png', '图 8  登录页：账号输入与注册、找回密码入口')
    figure(doc, '05-home-mobile.png', '图 9  移动端首页：导航折叠与单列文章列表', width=2.25)
    table(doc, ('界面元素', '输入 / 状态', '设计说明'), [
        ('全站导航', '桌面展开；窄屏折叠', '保留首页、主题栏目、归档、搜索及主题切换入口'),
        ('文章卡片', '标题、日期、摘要、标签', '标题进入详情；栏目与标签分别进入筛选列表'),
        ('文章详情', '目录、正文、推荐、评论', '长文通过目录定位；面包屑帮助返回上级'),
        ('评论区域', '未登录 / 已登录 / 审核状态', '未登录提示先登录；评论是否公开受 is_enable 与配置控制'),
        ('搜索', '空关键词 / 有结果 / 无结果', '关键词显示在结果标题，搜索依赖 Whoosh 或配置的 Elasticsearch 索引'),
        ('响应式与主题', '桌面 / 手机；浅色 / 深色', '小屏采用单列与移动导航；主题可由用户切换'),
    ])

    h2(doc, '模板设计')
    para(doc, '模板以两个共享布局为根：普通内容页继承 share_layout/base.html，登录与注册等账号页继承 share_layout/base_account.html。共享布局通过 include 组合导航和页脚；文章详情页通过 include 组合现代评论列表及发表评论组件。')
    figure(doc, '08-template-packages.png', '图 10  UML 包图：布局、页面、局部模板与自定义标签库的依赖关系')
    table(doc, ('关系', '代码实例', '作用'), [
        ('继承 extends', 'article_index.html / article_detail.html → share_layout/base.html', '复用页面骨架与 content 等 block'),
        ('继承 extends', 'account/login.html → share_layout/base_account.html', '账号页独立布局，不继承普通内容页'),
        ('包含 include', 'base.html → nav.html / footer.html', '导航与页脚按局部模板组合'),
        ('包含 include', 'article_detail.html → comment_list_modern.html / post_comment_modern.html', '文章详情中组合评论展示和登录后的评论输入'),
        ('标签 load / inclusion_tag', 'blog_tags.load_article_detail → blog/tags/article_info.html；load_sidebar → blog/tags/sidebar.html', '把 Python 计算结果注入小组件模板'),
        ('资源标签', 'vite_tags.vite_js / vite_preload', '从 Vite 清单加载前端资源'),
    ])
    para(doc, '说明：Django 模板的 extends 属于模板继承；include 为组合关系；自定义 inclusion_tag 是从标签函数到局部模板的渲染依赖，并非 Python 类继承。图示按此语义区分。')
    doc.save(OUT / 'AI Agent工程实践-软件界面设计说明书.docx')


def model_section(doc, name, description, rows):
    h3(doc, name)
    description_paragraph = para(doc, description)
    description_paragraph.paragraph_format.keep_with_next = True
    table(doc, ('字段', 'Django 字段类型', '约束 / 关系', '职责'), rows)


def build_data():
    doc = prepare(DATA_TEMPLATE, 'AI Agent 工程实践博客系统的数据模型设计', '软件数据模型设计说明书')
    h2(doc, '博客系统介绍')
    para(doc, '本项目是以 AI Agent 工程实践为内容主题的 Django 博客。数据模型服务于文章发布、栏目与标签组织、用户认证、评论交流和站点配置；它并非 Agent 任务执行平台，因此不存在 AgentTask、ToolCall 等运行实体。下文以当前 Django 模型和迁移文件为依据。')

    h2(doc, '系统数据分析')
    table(doc, ('业务对象', '需要保存的数据', '对应模型'), [
        ('用户与作者', '用户名、邮箱、密码哈希、昵称、来源', 'accounts.BlogUser'),
        ('文章内容', '标题、Markdown 正文、作者、栏目、状态、类型、发布时间、阅读量', 'blog.Article'),
        ('内容组织', '父子栏目、栏目顺序、标签及文章标签关联', 'blog.Category、blog.Tag、Article.tags'),
        ('读者互动', '评论正文、作者、所属文章、父评论、审核状态、表情反应', 'comments.Comment、CommentReaction'),
        ('站点配置', '站点名称、描述、关键词、色彩方案、评论开关', 'blog.BlogSettings'),
        ('扩展数据', '友情链接、侧栏、自第三方账号绑定及相关配置', 'blog.Links、SideBar；oauth.OAuthUser、OAuthConfig'),
    ])
    para(doc, '第四周独立演示数据库使用 utf8mb4，执行迁移后生成 12 篇主题文章、1 个根栏目及 5 个子栏目、14 个标签、3 条评论和 2 个不可登录的演示用户。数据命令可重复执行；不会删除已有数据，故在已有协作数据库运行时可能同时保留旧主题文章。')
    h3(doc, '数据规则与访问边界')
    table(doc, ('规则', '实现依据', '作用'), [
        ('文章标题、栏目名、标签名唯一', 'Article.title、Category.name、Tag.name 的 unique=True', '避免重复标识与导航歧义'),
        ('文章归属', 'Article.author/category 外键；tags 多对多', '文章须有作者与栏目，可关联多个标签'),
        ('评论树', 'Comment.parent_comment 自关联，可空', '支持一级评论与嵌套回复'),
        ('评论展示', 'Comment.is_enable 默认 False；站点和文章另有评论开关', '审核及关闭评论时不直接公开内容'),
        ('反应去重', 'CommentReaction(user, comment, reaction_type) 联合唯一', '同一用户对同一评论的同类反应只计一次'),
        ('级联删除', '核心外键 on_delete=CASCADE', '删除父对象将影响子对象，应在管理操作前确认'),
    ])

    h2(doc, '数据模型设计')
    figure(doc, '09-data-model-er.png', '图 1  核心数据模型 ER 图（辅助配置表未全部展开）')
    table(doc, ('关联', '基数', '实现字段 / 关联表'), [
        ('BlogUser → Article', '1 : N', 'Article.author_id'),
        ('Category → Article', '1 : N', 'Article.category_id'),
        ('Category → Category', '1 : N', 'Category.parent_category_id'),
        ('Article ↔ Tag', 'N : M', 'Article.tags；Django 自动表 blog_article_tags'),
        ('Article → Comment', '1 : N', 'Comment.article_id'),
        ('BlogUser → Comment', '1 : N', 'Comment.author_id'),
        ('Comment → Comment', '1 : N', 'Comment.parent_comment_id'),
        ('Comment → CommentReaction', '1 : N', 'CommentReaction.comment_id'),
    ])

    model_section(doc, 'BlogUser（accounts_bloguser）', '继承 Django AbstractUser，负责登录身份、作者与评论者归属。邮箱不是唯一登录标识；用户名继承唯一约束。', [
        ('id', 'AutoField', '主键', '用户标识'), ('username', 'CharField(150)', '唯一', '登录名 / 个人页标识'),
        ('email', 'EmailField(254)', '可重复', '邮件与联系信息'), ('password', 'CharField(128)', '密码哈希', '认证凭据'),
        ('nickname', 'CharField(100)', '可空字符串', '显示昵称'), ('source', 'CharField(100)', '可空字符串', '注册来源'),
    ])
    model_section(doc, 'Category（blog_category）', '栏目采用自关联树；本周配置一个 AI Agent 根栏目和五个子栏目。', [
        ('id', 'AutoField', '主键', '栏目标识'), ('name', 'CharField(30)', '唯一', '栏目名称'),
        ('parent_category', 'ForeignKey(self)', '可空，CASCADE', '上级栏目'), ('slug', 'SlugField(60)', '保存时生成', '栏目 URL 标识'),
        ('index', 'IntegerField', '默认 0', '导航排序'),
    ])
    model_section(doc, 'Tag（blog_tag）', '标签用于跨栏目的主题聚合，如 RAG、Safety、Workflow。', [
        ('id', 'AutoField', '主键', '标签标识'), ('name', 'CharField(30)', '唯一', '标签名称'),
        ('slug', 'SlugField(60)', '保存时生成', '标签 URL 标识'),
    ])
    model_section(doc, 'Article（blog_article）', '核心内容实体。文章在发布后进入首页、栏目、标签、归档与搜索结果；草稿不应进入公开列表。', [
        ('id', 'AutoField', '主键', '文章标识'), ('title', 'CharField(200)', '唯一', '文章标题'),
        ('body', 'MDTextField', '必填', 'Markdown 正文'), ('author', 'ForeignKey(BlogUser)', '必填，CASCADE', '作者'),
        ('category', 'ForeignKey(Category)', '必填，CASCADE', '所属栏目'), ('tags', 'ManyToManyField(Tag)', '可空', '主题标签'),
        ('status', 'CharField(1)', 'd 草稿 / p 发布', '公开状态'), ('type', 'CharField(1)', 'a 文章 / p 独立页', '内容类型'),
        ('comment_status', 'CharField(1)', 'o 开启 / c 关闭', '文章评论开关'), ('pub_time', 'DateTimeField', '非空', '发布时间'),
        ('views', 'PositiveIntegerField', '默认 0', '阅读计数'), ('show_toc', 'BooleanField', '默认 False', '是否展示目录'),
    ])
    model_section(doc, 'Comment（comments_comment）', '读者对文章的评论及回复。公开列表由 is_enable 控制。', [
        ('id', 'AutoField', '主键', '评论标识'), ('body', 'TextField(max_length=300)', '模型层长度限制', '评论内容'),
        ('article', 'ForeignKey(Article)', '必填，CASCADE', '所属文章'), ('author', 'ForeignKey(BlogUser)', '必填，CASCADE', '评论人'),
        ('parent_comment', 'ForeignKey(self)', '可空，CASCADE', '被回复评论'), ('is_enable', 'BooleanField', '默认 False', '审核 / 可见状态'),
        ('creation_time', 'DateTimeField', '默认当前时间', '提交时间'),
    ])
    doc.add_page_break()
    model_section(doc, 'CommentReaction（comments_commentreaction）', '评论的表情反应。', [
        ('id', 'AutoField', '主键', '反应标识'), ('comment', 'ForeignKey(Comment)', '必填，CASCADE', '所属评论'),
        ('user', 'ForeignKey(BlogUser)', '必填，CASCADE', '操作用户'), ('reaction_type', 'CharField(10)', '枚举值', '表情类型'),
        ('created_at', 'DateTimeField', 'auto_now_add', '反应时间'),
    ])
    model_section(doc, 'BlogSettings（blog_blogsettings）', '站点级配置逻辑上只使用一条记录，clean() 会检查重复配置；这不是数据库唯一约束。', [
        ('id', 'AutoField', '主键', '配置记录'), ('site_name', 'CharField(200)', '非空', '站点标题'),
        ('site_description', 'TextField(max_length=1000)', '非空', '站点简介'), ('site_keywords', 'TextField(max_length=1000)', '非空', 'SEO 关键词'),
        ('color_scheme', 'CharField(20)', '预设主题枚举', '界面配色'), ('open_site_comment', 'BooleanField', '默认 True', '全站评论开关'),
        ('comment_need_review', 'BooleanField', '默认 False', '评论审核策略'),
    ])
    h3(doc, '索引与扩展模型')
    para(doc, 'Article 已定义 type+status+pub_time、status+views、author+status+type、category+status 组合索引；Comment 已定义 article+parent_comment+is_enable 与 is_enable+id 索引；CommentReaction 对 comment+reaction_type 建索引。Links、SideBar、OAuthUser、OAuthConfig 等为友情链接、侧栏与第三方认证提供扩展，不改变图 1 的核心内容链路。')
    h3(doc, '初始化和复核步骤')
    para(doc, '在全新 MySQL 数据库设置 DJANGO_MYSQL_DATABASE 与凭据后，执行 python manage.py migrate、python manage.py create_testdata；搜索功能再执行 python manage.py update_index。不要将数据库密码写入仓库，不要在已有协作数据库直接清空旧内容。回归测试：python manage.py test blog.test_agent_demo_data --noinput。')
    doc.save(OUT / 'AI Agent工程实践-软件数据模型设计说明书.docx')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    build_ui()
    build_data()
    print('Created UI and data-model DOCX files in', OUT)
