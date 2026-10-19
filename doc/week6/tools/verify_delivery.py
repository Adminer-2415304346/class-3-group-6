"""验证用例覆盖、DFD平衡、注释行为一致性及交付文档结构。"""
from pathlib import Path
import ast
import hashlib
import io
import json
import re
import subprocess
import tokenize
import zipfile
import xml.etree.ElementTree as ET
from docx import Document

WEEK=Path(__file__).resolve().parents[1]
ROOT=WEEK.parents[1]

def main():
    cases=json.loads((WEEK/'content/usecases.json').read_text(encoding='utf-8'))['cases']
    data=json.loads((WEEK/'content/architecture.json').read_text(encoding='utf-8'))
    expected={f'UC{i:02}' for i in range(1,11)}
    assert {c['id'] for c in cases}==expected and len(cases)==10
    messages=0
    for c in cases:
        for k in ['name','actors','goal','trigger','preconditions','main_flow','alternatives','postconditions','source_refs','participants','messages']:
            assert c[k],(c['id'],k)
        ids={p['id'] for p in c['participants']}
        assert len(ids)==len(c['participants'])
        for m in c['messages']:
            assert m['from'] in ids and m['to'] in ids
            assert m['type'] in ['call','return']
        messages+=len(c['messages'])
        assert (WEEK/'assets'/(c['id']+'-顺序图.png')).is_file()
        assert (WEEK/'图源'/(c['id']+'-顺序图.puml')).is_file()
    dd=data['dictionary']
    flows={x['id']:x for x in dd['flows']}
    structures={x['id'] for x in dd['structures']}
    for gname in ['context','level1']:
        graph=data[gname];nodes={x['id']:x for x in graph['nodes']}
        assert len(nodes)==len(graph['nodes'])
        assert len({e['id'] for e in graph['edges']})==len(graph['edges'])
        for e in graph['edges']:
            assert e['from'] in nodes and e['to'] in nodes
            assert e['id'] in flows and e['structure'] in structures
            assert 'process' in (nodes[e['from']]['kind'],nodes[e['to']]['kind'])
    def external_signature(graph):
        ext={n['id'] for n in graph['nodes'] if n['kind']=='entity'}
        return {(e['id'],e['from'] if e['from'] in ext else e['to'],'in' if e['from'] in ext else 'out',e['structure']) for e in graph['edges'] if e['from'] in ext or e['to'] in ext}
    assert external_signature(data['context'])==external_signature(data['level1'])
    source='src/DjangoBlog/blog/views.py'
    old=subprocess.check_output(['git','show','HEAD:'+source],cwd=ROOT).decode('utf-8')
    new=(ROOT/source).read_text(encoding='utf-8')
    assert ast.dump(ast.parse(old),include_attributes=False)==ast.dump(ast.parse(new),include_attributes=False)
    def noncomments(s):
        ignored={tokenize.COMMENT,tokenize.NL,tokenize.NEWLINE,tokenize.ENCODING,tokenize.ENDMARKER}
        return [(t.type,t.string) for t in tokenize.generate_tokens(io.StringIO(s).readline) if t.type not in ignored]
    assert noncomments(old)==noncomments(new)
    assert new.count('# SZX（第六周）')==3
    docx=WEEK/'AI Agent工程实践-开源软件泛读报告.docx'
    d=Document(docx)
    text='\n'.join(p.text for p in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
    for word in ['1 开源软件的功能描述','2 界面分析','3 数据模型分析','4 软件架构及各个包和类的作用','5 软件功能与类间的对应关系','6 总结']:
        assert word in text
    for uid in expected:assert text.count(uid)>=4
    assert '***' not in text and '组号：xx' not in text
    with zipfile.ZipFile(docx) as z:
        assert z.testzip() is None
        for name in z.namelist():
            if name.endswith('.xml'):ET.fromstring(z.read(name))
    pages=ET.parse(WEEK/'图源/系统分析模型.drawio').getroot().findall('diagram')
    assert len(pages)==17,len(pages) # 用例、架构、上下文、4个一层分区、10顺序图
    assert not list(WEEK.glob('*.pdf')), '本轮按用户要求仅交付Word版本'
    result={'usecases':len(cases),'sequence_messages':messages,'diagrams':len(pages),'flows':len(flows),'structures':len(structures),'elements':len(dd['elements']),'stores':len(dd['stores']),'external_flows_balanced':len(external_signature(data['context'])),'source_ast_unchanged':True,'source_noncomment_tokens_unchanged':True,'new_annotation_groups':3,'docx_images':len(d.inline_shapes),'docx_tables':len(d.tables),'report_format':'docx'}
    (WEEK/'content/核验结果.json').write_text(json.dumps(result,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(result,ensure_ascii=True,indent=2))

if __name__=='__main__':main()
