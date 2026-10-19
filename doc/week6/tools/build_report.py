"""按课程泛读报告模板构建第六周报告；输入内容均由源码分析所得。

依赖 python-docx、Pillow。不会导入 Django、连接数据库或调用外部服务。
"""
from pathlib import Path
import json
import subprocess
import xml.etree.ElementTree as ET
from copy import deepcopy
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from diagram_engine import Canvas, sequence, drawio, wrap, BLUE, GRAY

WEEK=Path(__file__).resolve().parents[1]
ROOT=WEEK.parents[1]
ASSETS=WEEK/'assets'
CONTENT=WEEK/'content'
OUT=WEEK/'AI Agent工程实践-开源软件泛读报告.docx'
TEMPLATE=(WEEK/'模板/开源软件泛读报告.docx') if (WEEK/'模板/开源软件泛读报告.docx').exists() else ROOT.parent/'文档模板-开源软件泛读报告.docx'
PAGES=[]

def read(name): return json.loads((CONTENT/name).read_text(encoding='utf-8'))

def runfont(r,size=10.5,bold=False,name='宋体'):
    r.font.name='Times New Roman'; r.font.size=Pt(size);r.bold=bold
    r.font.color.rgb=RGBColor.from_string('1D2E47')
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),name)

def normalize_property_order(d):
    # Word可容错显示的属性顺序仍需符合OOXML，避免其它阅读器解析差异。
    orders={
      'pPr':'pStyle keepNext keepLines pageBreakBefore framePr widowControl numPr suppressLineNumbers pBdr shd tabs suppressAutoHyphens kinsoku wordWrap overflowPunct topLinePunct autoSpaceDE autoSpaceDN bidi adjustRightInd snapToGrid spacing ind contextualSpacing mirrorIndents suppressOverlap jc textDirection textAlignment textboxTightWrap outlineLvl divId cnfStyle rPr sectPr pPrChange'.split(),
      'tcPr':'cnfStyle tcW gridSpan hMerge vMerge tcBorders shd noWrap tcMar textDirection tcFitText vAlign hideMark headers cellIns cellDel cellMerge tcPrChange'.split()}
    for name,order in orders.items():
        rank={qn('w:'+x):i for i,x in enumerate(order)}
        for pr in d._element.body.iter(qn('w:'+name)):
            children=list(pr)
            children.sort(key=lambda el:rank.get(el.tag,len(order)))
            for el in children:pr.append(el)

def body(d,text,small=False):
    p=d.add_paragraph()
    p.paragraph_format.space_after=Pt(5)
    p.paragraph_format.line_spacing=1.15
    nogrid=OxmlElement('w:snapToGrid');nogrid.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(nogrid)
    runfont(p.add_run(str(text)),9 if small else 10.5)
    return p

def heading(d,text,level=2,page=False):
    p=d.add_heading(text,level)
    p.paragraph_format.page_break_before=page
    p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_before=Pt(10)
    p.paragraph_format.space_after=Pt(6)
    for r in p.runs: runfont(r,14 if level==1 else 11.5,True,'微软雅黑')
    return p

def table(d,headers,rows,widths=None,size=9):
    t=d.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    widths=widths or [16.6/len(headers)]*len(headers)
    for col,w in zip(t.columns,widths):col.width=Cm(w)
    for i,values in enumerate([headers]+list(rows)):
        row=t.rows[0] if i==0 else t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell,value,w in zip(row.cells,values,widths):
            cell.width=Cm(w);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            mar=OxmlElement('w:tcMar')
            for key in ['top','left','bottom','right']:
                el=OxmlElement('w:'+key);el.set(qn('w:w'),'65');el.set(qn('w:type'),'dxa');mar.append(el)
            tcpr.append(mar)
            borders=OxmlElement('w:tcBorders')
            for key in ['top','left','bottom','right']:
                el=OxmlElement('w:'+key);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'CBD5E1');borders.append(el)
            tcpr.append(borders)
            if i==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:val'),'clear');sh.set(qn('w:fill'),'EAF0F7');tcpr.append(sh)
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(1);p.paragraph_format.line_spacing=1.05
            p.paragraph_format.keep_with_next=False
            nogrid=OxmlElement('w:snapToGrid');nogrid.set(qn('w:val'),'0');p._p.get_or_add_pPr().append(nogrid)
            runfont(p.add_run(string(value)),size,i==0)
    spacer=body(d,'',True)
    spacer.paragraph_format.space_after=Pt(0);spacer.paragraph_format.line_spacing=Pt(1)
    for r in spacer.runs:r.font.size=Pt(1)
    return t

def string(v):
    if v is None:return ''
    if isinstance(v,list):return '\n'.join(string(x) for x in v)
    if isinstance(v,dict):return '；'.join(str(k)+'：'+string(x) for k,x in v.items())
    return str(v)

def figure(d,path,caption,max_height=21.5):
    from PIL import Image
    path=Path(path)
    w,h=Image.open(path).size
    width=min(16.6,max_height*w/h)
    p=d.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after=Pt(3)
    p.add_run().add_picture(str(path),width=Cm(width))
    p._p.xpath('.//wp:docPr')[0].set('descr',caption)
    p=body(d,caption,True);p.alignment=WD_ALIGN_PARAGRAPH.CENTER

def usecase_diagram(cases):
    c=Canvas('系统用例图：DjangoBlog 技术博客',2400,1740,'UC01—UC10 与用例说明、顺序图、功能—类映射使用同一编号。')
    c.shape((420,170,1560,1480),'',fill='#FFFFFF')
    c.text(1200,195,'AI Agent 工程实践博客系统',29,center=True,bold=True)
    centers={}
    # 左侧是公开与账号入口，右侧是登录互动与管理。
    for i,u in enumerate(cases):
        if i<6:x,y=790,370+i*205
        else:x,y=1605,440+(i-6)*315
        centers[u['id']]=(x,y)
        c.shape((x-285,y-69,570,138),wrap(u['id']+' '+u['name'],13),'ellipse',size=36)
    actors=[(160,670,'访客'),(2250,550,'读者'),(2250,1010,'第三方账号用户'),(2250,1470,'后台管理员')]
    for x,y,label in actors:
        c.shape((x-65,y-80,130,90),'','ellipse',fill='white')
        c.line([(x,y+10),(x,y+110)]);c.line([(x-80,y+55),(x+80,y+55)])
        c.line([(x-70,y+170),(x,y+110),(x+70,y+170)])
        c.text(x,y+185,label,29,center=True)
    for i,u in enumerate(cases):
        x,y=centers[u['id']]
        if i<6:pts=[(240,720),(505,y)]
        elif i<8:pts=[(2170,610),(1890,y)]
        elif i==8:pts=[(2170,1070),(1890,y)]
        else:pts=[(2170,1520),(1890,y)]
        c.line(pts,color=GRAY,width=2)
        if i==6:c.text(2060,440,'需登录',24,center=True,bg=True)
        if i==7:c.text(2110,755,'GET公开\nPOST登录',23,center=True,bg=True)
    c.text(70,1680,'已登录用户也可访问公开内容；当前自定义后台以 is_superuser 判断入口权限。OAuth 与邮件在数据流图中展开。',25,color=GRAY)
    c.save(ASSETS,'01-用例图');PAGES.append(c)

def package_diagram():
    c=Canvas('系统架构图：Django MVT 与应用包依赖',2200,1680,'虚线箭头表示依赖；前端通过 HTTP 与模板约定协作，业务持久化统一经过 Django ORM。')
    boxes={
      'browser':(710,145,780,125,'浏览器 / frontend\nVite · Alpine/htmx · CSS'),
      'templates':(120,400,520,155,'«package» templates\n继承、包含、模板标签'),
      'core':(840,400,520,155,'«package» djangoblog\nURLConf · 中间件 · 共享基类'),
      'blog':(120,755,440,155,'«package» blog\n内容、列表、详情、搜索'),
      'accounts':(625,755,440,155,'«package» accounts\n注册、激活、会话、重置'),
      'comments':(1130,755,440,155,'«package» comments\n评论、回复、表情'),
      'oauth':(1635,755,440,155,'«package» oauth\n渠道配置与外部身份'),
      'model':(320,1110,780,150,'«package» models / Django ORM\nArticle · BlogUser · Comment · OAuthUser'),
      'service':(1290,1110,780,150,'«package» blog_signals / plugins\n信号邮件处理 · 通知 · Hook 扩展'),
      'database':(320,1450,780,125,'数据库 / 缓存 / 搜索索引\n关系存储 · Django cache · Haystack/ES'),
      'external':(1290,1450,780,125,'外部系统（配置后启用）\nOAuth 平台 · SMTP · 百度 / Cloudflare')}
    routes=[('browser','core',[(1100,270),(1100,400)]),('core','templates',[(840,477),(640,477)]),
      ('core','blog',[(960,555),(960,635),(340,635),(340,755)]),('core','accounts',[(1030,555),(1030,680),(845,680),(845,755)]),
      ('core','comments',[(1170,555),(1170,680),(1350,680),(1350,755)]),('core','oauth',[(1250,555),(1250,635),(1855,635),(1855,755)]),
      ('blog','model',[(340,910),(340,1010),(570,1010),(570,1110)]),('accounts','model',[(845,910),(845,1110)]),
      ('comments','model',[(1350,910),(1350,1040),(1000,1040),(1000,1110)]),('oauth','model',[(1855,910),(1855,1000),(1090,1000),(1090,1110)]),
      ('model','database',[(710,1260),(710,1450)]),('blog','service',[(540,910),(540,960),(1590,960),(1590,1110)]),
      ('accounts','service',[(1000,910),(1000,990),(1460,990),(1460,1110)]),('service','external',[(1680,1260),(1680,1450)]),
      ('oauth','external',[(2075,832),(2140,832),(2140,1510),(2070,1510)])]
    for a,b,pts in routes:c.line(pts,dashed=True,arrow=True,color=GRAY)
    for key,(x,y,w,h,label) in boxes.items():c.shape((x,y,w,h),label,'package' if key not in ['browser','database','external'] else 'rect',size=33)
    c.save(ASSETS,'04-系统架构图');PAGES.append(c)

def graph_data(data,stem,title):
    """将经核对的数据流模型绘成布局固定的分区图；边以编号索引字典。"""
    # 模型的 x/y/w/h 可由编制者指定；没有时使用层次布局。
    nodes=data['nodes'];edges=data['edges']
    is_context='context' in stem
    width,height=data.get('width',2400),data.get('height',1600 if is_context else 1900)
    c=Canvas(title,width,height,'矩形：外部实体    圆形/椭圆：处理    开口框：数据存储    箭头：数据流（编号对应字典）')
    boxes={}
    procs=[n for n in nodes if n.get('kind') in ['process','system']]
    ext=[n for n in nodes if n not in procs and n.get('kind') not in ['store','data_store','datastore']]
    stores=[n for n in nodes if n.get('kind') in ['store','data_store','datastore']]
    for i,n in enumerate(procs):
        if is_context: box=(900,650,600,230)
        else:box=(250+i*550,650,440,175)
        boxes[n['id']]=box
    for i,n in enumerate(ext):
        if is_context:
            pos=[(60,270),(60,1120),(1840,270),(1840,1120),(900,1370)]
            x,y=pos[i%len(pos)]; box=(x,y,500,130)
        else:
            box=(80+i*560,230,430,125)
        boxes[n['id']]=box
    for i,n in enumerate(stores):
        row=i//4;col=i%4
        boxes[n['id']]=(130+col*580,1170+row*320,430,145)
    # 指定的几何只覆盖默认值，便于手工微调。
    for n in nodes:
        if all(k in n for k in ('x','y','w','h')):boxes[n['id']]=tuple(n[k] for k in ('x','y','w','h'))
    if is_context:boxes['P0']=(720,580,650,250)
    pairs={}
    for e in edges:
        key=tuple(sorted([e['from'],e['to']]))
        pairs.setdefault(key,[]).append(e)
    for pair,es in pairs.items():
        directed={}
        for e in es:directed.setdefault((e['from'],e['to']),[]).append(e)
        for j,((src,dst),bundle) in enumerate(directed.items()):
            e=bundle[0]
            a=boxes[e['from']];b=boxes[e['to']]
            ac=(a[0]+a[2]/2,a[1]+a[3]/2);bc=(b[0]+b[2]/2,b[1]+b[3]/2)
            off=(j-(len(directed)-1)/2)*65
            if abs(ac[0]-bc[0])>abs(ac[1]-bc[1]):
                sign=1 if bc[0]>ac[0] else -1
                s=(ac[0]+sign*a[2]/2,ac[1]+off);t=(bc[0]-sign*b[2]/2,bc[1]+off)
                mid=(s[0]+t[0])/2+off
                pts=[s,(mid,s[1]),(mid,t[1]),t]
                lp=(mid+off,(s[1]+t[1])/2-35)
            else:
                sign=1 if bc[1]>ac[1] else -1
                s=(ac[0]+off,ac[1]+sign*a[3]/2);t=(bc[0]+off,bc[1]-sign*b[3]/2)
                mid=(s[1]+t[1])/2+off
                pts=[s,(s[0],mid),(t[0],mid),t]
                lp=((s[0]+t[0])/2,mid-36)
            # 单独一层图只放编号，避免把数据字典搬进图中。
            label='/'.join(x['id'] for x in bundle)
            c.line(pts,wrap(label,14),arrow=True,color=GRAY,width=2,labelpos=lp)
    for n in nodes:
        k=n.get('kind','entity')
        shape='ellipse' if k in ['system','process'] else 'store' if k in ['store','data_store','datastore'] else 'rect'
        label=node_label(n)
        if is_context and n['id']=='P0':label='P0 DjangoBlog\n核心 Web 系统'
        c.shape(boxes[n['id']],wrap(label,12),shape,size=34)
    c.text(60,height-55,'同向并行数据流合并标注全部编号，定义见字典；相同外部输入/输出在上下文图与一层图中保持平衡。',24,color=GRAY)
    c.save(ASSETS,stem);PAGES.append(c)

def node_label(n):
    import re
    label=n['label']
    label=re.sub(r'^'+n['id']+r'\s*','',label)
    label=re.sub(r'^\d\s+','',label)
    if n['id']=='E5':label='百度 / Cloudflare\nURL 服务'
    return n['id']+' '+label

def level1_panels(data):
    """同一层DFD跨四张纸分区，重复节点仍用原编号，不引入新的分解层。"""
    byid={n['id']:n for n in data['nodes']}
    for pid in ['P1','P2','P3','P4']:
        edges=[e for e in data['edges'] if pid in (e['from'],e['to'])]
        neighbor_ids=list(dict.fromkeys(e['to'] if e['from']==pid else e['from'] for e in edges))
        left=[byid[x] for x in neighbor_ids if byid[x]['kind']!='store']
        right=[byid[x] for x in neighbor_ids if byid[x]['kind']=='store']
        height=max(1400,320+max(len(left),len(right))*230)
        c=Canvas(f'一层数据流图 {pid} 分区',2400,height,'四个分区共同组成同一层模型；重复外部实体、处理和存储使用同一编号，完整定义见数据字典。')
        center=(900,height/2-90,600,200)
        boxes={pid:center}
        for group,x in [(left,80),(right,1800)]:
            for i,n in enumerate(group):
                y=200+(i+.5)*(height-380)/max(len(group),1)-75
                boxes[n['id']]=(x,y,500,150)
        bundles={}
        for e in edges:bundles.setdefault((e['from'],e['to']),[]).append(e)
        for (src,dst),flows in bundles.items():
            other=dst if src==pid else src
            box=boxes[other]
            isleft=box[0]<900
            # 每对节点的去/回流采用不同端口，标签靠近外侧节点以避开汇聚区。
            outgoing=src==pid;off=-32 if outgoing else 32
            side=900 if isleft else 1500
            near=box[0]+500 if isleft else box[0]
            cy=height/2+10+off;ny=box[1]+75+off
            pts=[(side,cy),(near,ny)] if outgoing else [(near,ny),(side,cy)]
            c.line(pts,arrow=True,color=GRAY,width=3)
            lx=near+110 if isleft else near-110
            # 标签用白底；同一端口的多个流逐行列出。
            label='/'.join(e['id'] for e in flows)
            c.text(lx,ny-39,wrap(label,11),30,center=True,bg=True)
        for nid,box in boxes.items():
            n=byid[nid]
            typ='store' if n['kind']=='store' else 'ellipse' if n['kind']=='process' else 'rect'
            c.shape(box,wrap(node_label(n),12),typ,size=34)
        c.text(60,height-55,'矩形：外部实体    椭圆：处理    开口框：数据存储    每条有向流均保留原F编号；同向多流合并标注。',26,color=GRAY)
        c.save(ASSETS,'03-level1-'+pid);PAGES.append(c)

def dict_table(d,title,items,keys=None):
    heading(d,title,2)
    if not items:return
    if keys is None:
        keys=list(dict.fromkeys(k for it in items for k in it))
    captions={'id':'编号','name':'名称','label':'名称','type':'类型','definition':'定义','structure':'结构','description':'说明','source':'来源','target':'去向','from':'来源','to':'去向','rules':'规则','fields':'组成','constraint':'约束','constraints':'约束','notes':'说明','model':'模型','models':'关联模型','path':'源码','paths':'源码路径','purpose':'用途','code':'源码','responsibility':'职责','evidence':'依据','key':'键与约束'}
    # 限制列数，额外属性合并进说明列，避免横向拥挤。
    if len(keys)>4:
        front=keys[:3];rest=keys[3:]
        rows=[[string(it.get(k,'')) for k in front]+['\n'.join(captions.get(k,k)+'：'+string(it[k]) for k in rest if k in it)] for it in items]
        table(d,[captions.get(k,k) for k in front]+['定义与规则'],rows,[1.5,3.0,4.0,8.1] if len(front)==3 else None,8.5)
    else:
        table(d,[captions.get(k,k) for k in keys],[[string(it.get(k,'')) for k in keys] for it in items],size=8.7)

def make_report(cases,a):
    d=Document(TEMPLATE)
    for el in list(d._element.body):
        if el.tag!=qn('w:sectPr'):d._element.body.remove(el)
    sec=d.sections[0]
    sec.top_margin=Cm(1.8);sec.bottom_margin=Cm(1.7)
    sec.left_margin=sec.right_margin=Cm(2.2)
    for name in ['Normal','Heading 1','Heading 2','Heading 3','Caption']:
        if name not in d.styles:continue
        st=d.styles[name];st.font.name='Times New Roman';st.font.size=Pt(10.5)
        st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
        st.paragraph_format.line_spacing=1.15
    p=d.add_paragraph();p.paragraph_format.space_before=Pt(65);p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    runfont(p.add_run('“AI Agent 工程实践”\n开源博客系统的泛读报告'),23,True,'微软雅黑')
    p=body(d,'软件工程方法学 · 大作业第六周');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    body(d,'')
    table(d,['项目','信息'],[
      ('组号','24级软件工程3班 第6组'),
      ('小组成员','石兆翔 2415304346；喻飞扬 2415304347；王子扬 2415304348\n顾梓鑫 2415304345；宋洋 2415304330；郭曹钰 2415304331'),
      ('项目','DjangoBlog 二次开发；内容主题为 AI Agent 工程实践'),
      ('编制与核对','石兆翔负责汇总；本报告为源码静态分析成果'),
      ('编制日期','2026年10月19日'),
      ('源码基线',subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip()),
      ('模板','《文档模板-开源软件泛读报告.docx》；保留六项正文结构，增加完整用例与顺序图附件')
    ],[3,13.6],10)
    body(d,'本报告分析已实现的技术博客。文章中的 Agent 案例属于内容主题，系统未因此具有自主智能体执行能力。图中控制类与边界类用于分析职责，不表示本周新增同名源码类。')
    body(d,'本次新完成内容：功能用例、上下文与一层数据流、主要数据字典、架构依赖、功能—类映射及各用例顺序图。界面截图与 ER 图承接第四、五周成果，并注明来源。',True)
    body(d,'本周实际组会日期、个人阅读过程与成员确认不由源码推断，见随附计划与学习记录。',True)
    # 页眉页码。
    hp=sec.header.paragraphs[0];hp.text='软件工程方法学  /  第6组  /  第六周泛读报告'
    for r in hp.runs:runfont(r,8)
    fp=sec.footer.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');fp._p.append(f)
    heading(d,'阅读与提交导航',1,True)
    table(d,['课程要求','本报告位置','配套成果'],[
      ('用例图、用例说明、每个用例顺序图','1.1、附录A（10个用例）、附录B（10张顺序图）','assets/01-用例图；图源/系统分析模型.drawio'),
      ('数据流图与主要数据字典','1.2—1.3','上下文图、一层图；content/architecture.json'),
      ('界面与数据模型','第2、3节','引用week4界面截图、week5 ER与字段说明'),
      ('系统架构、包和类职责','第4、5节','UML包图、职责与功能追踪表'),
      ('前阶段总结、本周计划与组会','第6节；独立计划文件','第六周计划与组会议程.md'),
      ('TSPi 6—9章、Django熟悉与标注','配套个人材料','TSPi学习辅助、自测、Django代码标注记录')
    ],[4.1,6.3,6.2])
    body(d,'分析方法：先沿 URLConf → 视图/表单 → 模型/共享组件 → 模板/响应追踪真实调用，再以业务职责建立用例、数据流和包依赖。带src/的路径相对仓库根目录，其余应用文件路径相对src/DjangoBlog；图的编号是本周统一编号。')
    body(d,'建模口径：用例图描述参与者目标；数据流图只表达数据输入、处理和存储；顺序图表达对象交互与响应先后；架构图表达包依赖。四类图不混用箭头含义。')
    heading(d,'1 开源软件的功能描述',1,True)
    body(d,'访客可阅读、筛选与搜索文章，进行注册、登录或密码找回；登录用户可评论、回复和进行表情互动；后台管理员维护文章、分类、标签、评论及站点配置。自定义后台入口实际检查 is_superuser，不能仅凭 is_staff 推断可以访问。第三方登录需要渠道凭据及网络服务。')
    heading(d,'1.1 用例图与系统边界',2)
    figure(d,ASSETS/'01-用例图.png','图1  系统用例图（完整说明见附录A）',18)
    heading(d,'1.1.1 用例目录与覆盖范围',2,True)
    table(d,['编号','用例','参与者','目标与范围'],[(u['id'],u['name'],'、'.join(u['actors']),u['goal']) for u in cases],[1.3,3.3,3,9],9)
    body(d,'UC10 的子任务包括内容编辑、发布/草稿切换、评论启用/禁用、站点开关、链接与侧边栏维护。同一后台交互模式在一个用例内列出备选流程，避免把后台每个字段拆成独立用例。辅助端点如 health、sitemap、RSS、上传、缓存维护列为实现支撑，不作为本次主要业务用例重复编号。')
    heading(d,'1.2 系统数据流图',2,True)
    figure(d,ASSETS/'02-context-顶层数据流图.png','图2  系统上下文数据流图',20)
    body(d,'上下文图把整个博客视为一个处理，所有外部实体通过数据流与系统交换数据。数据库、索引和缓存位于系统数据处理边界内，在一层图展开。',True)
    heading(d,'1.2.1 一层数据流图',2,True)
    body(d,'P1/P2/P3/P4 分别对应内容、账号、互动、后台处理。为清晰排版，将同一层模型分成四个分区；重复实体、处理和存储使用同一编号，不表示新增数据副本或下一层分解。公共外部流保持上下文图编号和含义。')
    for i,pid in enumerate(['P1','P2','P3','P4']):
        if i:heading(d,'1.2.1.'+str(i+1)+' '+pid+' 分区',2,True)
        figure(d,ASSETS/('03-level1-'+pid+'.png'),'图3-'+str(i+1)+'  一层数据流图 '+pid+' 分区（完整定义见数据字典）',19)
    body(d,'一层内部存储流不出现在上下文边界。搜索索引属于派生数据，缓存不是权威数据源。四个分区中重复出现的处理之间的流只计一次。',True)
    heading(d,'1.3 主要数据条目的数据字典',2,True)
    body(d,'记法：A+B 表示组成；{A} 表示重复；[A] 表示可选；A|B 表示择一。字典分为数据流、数据结构、数据元素与存储四类；枚举、NULL、unique 与业务校验分别说明，不把文字需求当作数据库约束。')
    dd=a['dictionary']
    for title,key in [('1.3.1 数据流字典','flows'),('1.3.2 数据结构字典','structures'),('1.3.3 主要数据元素','elements'),('1.3.4 数据存储字典','stores')]:
        dict_table(d,title,dd[key])
    heading(d,'2 界面分析',1,True)
    body(d,'本节承接 doc/week4/AI Agent工程实践-软件界面设计说明书.docx 的页面分析。下列图为第四周已有运行截图，用于说明界面职责，不作为2026年10月19日重新运行截图。当前源码中的路由、模板与账号页面变化以本报告映射为准。')
    figure(d,ASSETS/'历史引用/01-home-desktop.png','图4  桌面首页：内容导航、文章列表、侧边栏（第四周截图）',12.3)
    table(d,['页面','主要元素与交互','源码与用例'],[
      ('首页/列表','文章摘要、分类/标签/作者导航、分页、主题切换','blog/article_index.html；UC01'),
      ('详情/评论','Markdown正文、目录、作者、相邻文章、评论与表情','blog/article_detail.html；UC02/07/08'),
      ('搜索','关键词输入、分页结果、异常与空结果','search/search.html；UC03'),
      ('账号页面','注册、登录、验证码/找回密码及结果反馈','account/*.html；UC04/05/06/09'),
      ('管理后台','模型列表、编辑表单、批处理动作','Django admin / blog/admin.py；UC10')
    ],[3,7.2,6.4])
    heading(d,'2.1 详情与界面流转',2,True)
    figure(d,ASSETS/'历史引用/02b-article-comments.png','图5  文章评论区域（第四周截图）',12)
    figure(d,ASSETS/'历史引用/07-ui-flow.png','图6  既有界面流转分析（第四周成果）',9)
    body(d,'评论区展示受文章 comment_status 与站点 OPEN_SITE_COMMENT 控制。前端隐藏入口不等于服务器端已对全部条件拒绝请求；提交视图的实际校验边界见第6节与 UC07。',True)
    heading(d,'3 数据模型分析',1,True)
    body(d,'本节承接第五周正式数据模型说明书。核心业务含11个模型、3个隐式多对多关联表；servermanager另有commands与EmailSendLog两个支撑模型。Django自带的权限、会话、管理日志表属于框架支撑。BaseModel为抽象模型，不生成业务表。AST清单含上述13个实体及抽象模型共14项。')
    figure(d,ASSETS/'历史引用/er-core.png','图7  核心内容、用户、评论与互动关系（第五周成果）',19)
    heading(d,'3.1 外围配置与约束',2,True)
    figure(d,ASSETS/'历史引用/er-configuration.png','图8  OAuth 与配置实体关系（第五周成果）',12)
    table(d,['规则','源码事实与含义'],[
      ('文章关系','Article.author、category 为必填外键；tags 为多对多。status d/p 与 comment_status o/c 是两个不同业务枚举。'),
      ('评论树','Comment.parent_comment 可空；顶级与回复共享表。is_enable 只有公开/不公开，无法单独区分待审与禁用。'),
      ('互动约束','CommentReaction 的 comment/user/reaction_type 联合唯一；同一用户可给同一评论使用不同表情。'),
      ('账号唯一性','username 为数据库唯一；注册表单校验邮箱重复不代表 email 列已有数据库唯一约束。'),
      ('OAuth关联','OAuthUser.author 可空；type/openid 查询不是联合唯一约束；OAuthConfig 与 OAuthUser 无外键。'),
      ('配置单例','BlogSettings.clean()/后台控制添加；未声明数据库单例约束，普通 save() 不自动执行 full_clean()。'),
      ('删除行为','外键多采用 Django CASCADE 收集策略；可空外键不因此变成 SET_NULL。')
    ],[3.4,13.2])
    body(d,'第五周文档已提供完整模型字段字典；本报告1.3节补充数据流与结构，使输入、处理、存储、输出和模型字段可互相追踪。密码仅以Django编码值保存；OAuth密钥、token和邮箱验证码不进入公开响应。',True)
    heading(d,'4 软件架构及各个包和类的作用',1,True)
    figure(d,ASSETS/'04-系统架构图.png','图9  系统架构与 UML 包依赖图',21)
    body(d,'Django 的 View 承担请求协调；Template 负责展示；Model 负责持久化与实体方法。分析层的 control 对应现有视图/表单等协作。虚线箭头表示使用依赖；models / ORM、blog_signals / plugins 是逻辑分组，实际目录与类见4.1节。',True)
    heading(d,'4.1 包、主要类与职责',2,True)
    dict_table(d,'模块职责清单',a['modules'])
    body(d,'请求主链：浏览器 → 项目 URLConf → 应用 URLConf → 视图/表单 → ORM与共享组件 → 模板或JSON响应。中间件在请求/响应边界协作；上下文处理器向相应模板注入站点数据；Hook 由业务调用点显式触发。')
    body(d,'外部依赖：搜索由Haystack与实际配置的后端协作；OAuth需要平台授权；utils.send_email触发信号，由blog_signals处理器同步调用SMTP并记录EmailSendLog，servermanager提供日志模型等支撑；Cloudflare/百度通知属于可配置扩展。源码存在接口不表示这些服务已部署或通过真实联调。')
    heading(d,'5 软件功能与类间的对应关系',1,True)
    table(d,['用例','视图/控制职责','实体/服务','源码入口与依据'],[(u['id']+' '+u['name'], '\n'.join(p['label'] for p in u['participants'] if p.get('kind')=='control'), '\n'.join(p['label'] for p in u['participants'] if p.get('kind') in ['entity','external']), '\n'.join(u['source_refs'][:3])) for u in cases],[3.2,4,4.1,5.3],8.7)
    body(d,'每个用例在附录A提供触发、前置、主流程、备选/异常和后置；附录B以相同编号给出交互顺序。后台通用类由 Django admin 提供，图中实体和框架对象是职责抽象。')
    heading(d,'6 总结',1,True)
    body(d,'前一阶段已建立个人分支到develop/main的协作路径，完成界面和数据模型分析，并对各成员代码标注做过集中审校。本周将分散的页面、模型与调用链整理成统一功能模型，使每个用例能追踪到参与类、数据流和源码。')
    body(d,'系统有评论审核支持：comment_need_review=True 时新评论保持不公开，可由后台启用；is_enable=False 本身不能区分“待审”与“已禁用”。单篇文章可通过 comment_status=c 关闭评论，后台提供开关批操作。站点 open_site_comment 参与模板入口展示。')
    findings=a.get('findings',[])
    for finding in findings:
        p=body(d,finding.get('id','')+' '+finding.get('title','')+'：'+finding.get('detail',string(finding)),True)
        p.paragraph_format.space_after=Pt(4)
    body(d,'下一步按缺陷影响和团队分工处理：先补文章详情与评论提交的发布/开关规则，再检查父评论归属、账号/OAuth约束和后台批处理后的缓存/索引一致性；修复应另建任务并配隔离回归。本次分析不擅自改业务规则。')
    body(d,'本周组会采用随附议程核对上述模型；各成员逐项说明本人负责模块并确认。真实日期、发言及个人TSPi阅读完成情况待本人记录；当前文档生成不作为这些线下活动已完成的证明。')
    heading(d,'附录A 用例说明',1,True)
    body(d,'所有用例共享系统时区及项目的路由/语言配置。前置条件仅描述本次选定场景，不替代源码已有的权限、校验或异常事实。与图中数据/对象名称同义的业务词可互相对应。')
    for i,u in enumerate(cases):
        heading(d,u['id']+' '+u['name'],2,page=i>0)
        table(d,['项目','说明'],[
          ('主要参与者','、'.join(u['actors'])),('目标与触发',u['goal']+'\n触发：'+u['trigger']),
          ('前置条件',u['preconditions']),('主要成功流程',[f'{n+1}. {s}' for n,s in enumerate(u['main_flow'])]),
          ('备选与异常',u['alternatives']),('后置条件',u['postconditions']),
          ('实现对应',u['source_refs']),('分析说明',u.get('notes',[]))
        ],[2.4,14.2],9.2)
        body(d,'对应顺序图：附录B '+u['id']+'。图源中提供同名PlantUML及可编辑diagrams.net页面。',True)
    for i,u in enumerate(cases):
        heading(d,('附录B 各用例顺序图\n' if i==0 else '')+u['id']+' '+u['name'],1 if i==0 else 2,True)
        scene='注册缺陷与独立邮件激活支路' if u['id']=='UC04' else '选定交互场景'
        figure(d,ASSETS/(u['id']+'-顺序图.png'),'图B'+str(i+1)+'  '+u['id']+' '+scene,21)
        body(d,'流程条件、异常分支和实现限制见附录A。对象为分析级职责抽象；外部服务消息表示接口调用，未声明完成实际网络测试。',True)
    heading(d,'附录C 来源与复核说明',1,True)
    table(d,['来源','采用内容'],[
      ('课程截图：大作业第6周任务','个人阅读TSPi 6—9章、熟悉Django与初步标注；小组计划、架构模型与泛读报告。'),
      ('课程泛读报告模板','功能、界面、数据模型、架构/包与类、功能—类映射、总结六项结构；模板要求的评论问题。'),
      ('当前仓库源码','accounts、blog、comments、oauth、servermanager、djangoblog、templates、frontend；源码路径见各用例与数据字典。'),
      ('doc/week4','界面说明书、既有运行截图、界面流转结果。'),
      ('doc/week5','数据模型说明书、ER图、前阶段计划及历史验收记录。历史测试不作为本周重跑结果。'),
      ('本周验证','见 第六周交付核验记录.md：图模型、DFD平衡、用例覆盖、DOCX结构与渲染检查。')
    ],[4.3,12.3])
    body(d,'可复现输入位于content/；图源位于图源/；tools/build_report.py基于本地模板重新生成。本报告不读取真实业务记录、密码或第三方凭据。',True)
    d.core_properties.title='AI Agent工程实践-开源软件泛读报告'
    d.core_properties.author='石兆翔 / 24级软件工程3班第6组'
    d.core_properties.subject='软件工程方法学第六周任务'
    normalize_property_order(d)
    d.save(OUT)
    return OUT

def plantuml(cases):
    folder=WEEK/'图源';folder.mkdir(exist_ok=True)
    for u in cases:
        lines=['@startuml','skinparam defaultFontName Microsoft YaHei',f'title {u["id"]} {u["name"]}']
        for p in u['participants']:
            kind=p.get('kind','participant');typ={'control':'control','entity':'entity','actor':'actor','boundary':'boundary','external':'participant'}.get(kind,'participant')
            lines.append(f'{typ} "{p["label"]}" as {p["id"]}')
        lines.append('autonumber')
        for m in u['messages']:
            arrow='-->' if m.get('type')=='return' else '->'
            lines.append(f'{m["from"]} {arrow} {m["to"]} : {m["text"]}')
        lines+=['@enduml','']
        (folder/(u['id']+'-顺序图.puml')).write_text('\n'.join(lines),encoding='utf-8')

def main():
    import shutil
    ASSETS.mkdir(exist_ok=True)
    history=ASSETS/'历史引用';history.mkdir(exist_ok=True)
    for stage,name in [('week4','01-home-desktop.png'),('week4','02b-article-comments.png'),('week4','07-ui-flow.png'),('week5','er-core.png'),('week5','er-configuration.png')]:
        if not (history/name).exists():shutil.copy2(ROOT/'doc'/stage/'assets'/name,history/name)
    (WEEK/'模板').mkdir(exist_ok=True)
    if not (WEEK/'模板/开源软件泛读报告.docx').exists():shutil.copy2(TEMPLATE,WEEK/'模板/开源软件泛读报告.docx')
    cases=read('usecases.json')['cases'];a=read('architecture.json')
    usecase_diagram(cases);package_diagram()
    graph_data(a['context'],'02-context-顶层数据流图','系统上下文数据流图')
    level1_panels(a['level1'])
    for u in cases:PAGES.append(sequence(u,ASSETS))
    plantuml(cases)
    drawio(PAGES,WEEK/'图源'/'系统分析模型.drawio')
    print(make_report(cases,a))

if __name__=='__main__':main()
