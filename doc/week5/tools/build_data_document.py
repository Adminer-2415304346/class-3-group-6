"""SZX：保留课程模板的版式与包部件，生成第五周数据模型说明书和可复查 ER 图。"""

import argparse
from copy import deepcopy
from datetime import datetime, timezone
from io import BytesIO
import json
from pathlib import Path
import sys
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[3]
WEEK = ROOT / 'doc' / 'week5'
ASSETS = WEEK / 'assets'
OUTPUT = WEEK / 'AI Agent工程实践-软件数据模型设计说明书.docx'
TEMPLATE = ROOT.parent / '文档模板-软件数据模型设计说明书模板.docx'
MEMBERS = [
    ('石兆翔', '2415304346'), ('喻飞扬', '2415304347'), ('王子扬', '2415304348'),
    ('顾梓鑫', '2415304345'), ('宋洋', '2415304330'), ('郭曹钰', '2415304331'),
]
MODEL_NAMES = {
    'accounts.BlogUser': ('用户', 'accounts/models.py', '继承 AbstractUser，统一账号、身份状态和权限；作为文章、评论及第三方账号的用户关联目标。'),
    'blog.Article': ('文章与独立页面', 'blog/models.py', '保存 Markdown 正文、发布状态、栏目、作者和展示排序；type 区分文章与独立页面。'),
    'blog.Category': ('分类', 'blog/models.py', '通过可空父分类组织层级栏目；分类名称唯一，slug 用于分类详情路由。'),
    'blog.Tag': ('标签', 'blog/models.py', '按关键词交叉组织文章，通过隐式关联表与文章形成多对多关系。'),
    'blog.Links': ('友情链接', 'blog/models.py', '维护外部链接、展示位置、顺序和是否启用。'),
    'blog.SideBar': ('自定义侧边栏', 'blog/models.py', '维护可展示的 HTML 内容、顺序及启用标志；不声明到其他实体的外键。'),
    'blog.BlogSettings': ('站点配置', 'blog/models.py', '保存站点元信息、展示数量、主题与评论策略；单配置要求由 clean() 校验，不是数据库唯一约束。'),
    'comments.Comment': ('评论与回复', 'comments/models.py', '关联文章和作者，通过可空父评论组织回复；is_enable 控制是否进入公开评论列表。'),
    'comments.CommentReaction': ('评论表情反应', 'comments/models.py', '记录用户对评论的表情反应，以评论、用户、表情三字段联合唯一约束防止同种反应重复。'),
    'oauth.OAuthUser': ('第三方账号绑定', 'oauth/models.py', '记录第三方身份信息，可关联本地用户；(type, openid) 用于业务查询，但没有联合唯一约束。'),
    'oauth.OAuthConfig': ('第三方登录配置', 'oauth/models.py', '记录登录渠道、应用凭据、回调地址和启用状态；同渠道配置由 clean() 校验，不是数据库唯一约束。'),
}
COMMON_FIELDS = {
    'id': '自增主键，定位实体记录。',
    'creation_time': '创建时间，缺省取当前时间。',
    'last_modify_time': '修改时间字段；缺省 now，不自动随更新改变。',
    'last_mod_time': '修改时间字段；缺省 now，不自动随更新改变。',
    'is_enable': '是否启用；具体可见性由调用方过滤。',
    'sequence': '展示顺序，值在本表中唯一。',
    'slug': '路由标识，保存时按名称生成；未设唯一。',
}
DESCRIPTIONS = {
    'accounts.BlogUser': {
        'password': 'Django 密码编码结果，不存储可读明文。',
        'last_login': '最近登录时间；尚未登录可为空。',
        'is_superuser': '是否超级用户；与后台权限判断相关。',
        'username': '登录用户名，在数据库中唯一。',
        'first_name': '继承的名字段，可在表单中留空。',
        'last_name': '继承的姓字段，可在表单中留空。',
        'email': '电子邮箱；模型未设置唯一约束。',
        'is_staff': '是否允许作为后台工作人员。',
        'is_active': '账号是否处于激活状态。',
        'date_joined': '继承的账号加入时间。',
        'nickname': '公开显示昵称，可留空。',
        'source': '账号创建来源，可留空。',
    },
    'blog.Article': {
        'title': '文章或独立页面标题，数据库唯一。', 'body': 'Markdown 正文，编辑器字段存储为长文本。',
        'pub_time': '发布时间，用于文章列表排序。', 'status': 'd=草稿；p=发布。',
        'comment_status': 'o=打开评论；c=关闭评论。', 'type': 'a=文章；p=独立页面。',
        'views': '浏览次数，初值 0，viewed() 增加。', 'author': '作者；引用 BlogUser。',
        'article_order': '文章排序权重，优先降序排列。', 'show_toc': '是否展示文章目录。',
        'category': '所属分类；引用 Category。',
    },
    'blog.Category': {'name': '分类名称，数据库唯一。', 'parent_category': '父分类；为空表示根分类。', 'index': '分类顺序权重，按降序排列。'},
    'blog.Tag': {'name': '标签名称，数据库唯一。'},
    'blog.Links': {'name': '链接显示名称，数据库唯一。', 'link': '外部链接 URL。', 'show_type': 'i=首页；l=列表；p=文章；a=全部；s=侧边栏。'},
    'blog.SideBar': {'name': '侧边栏标题。', 'content': '侧边栏展示的 HTML 内容。'},
    'blog.BlogSettings': {
        'site_name': '站点名称。', 'site_description': '站点介绍文本。', 'site_seo_description': '搜索引擎描述文本。',
        'site_keywords': '站点关键词。', 'article_sub_length': '列表摘要截取长度。',
        'sidebar_article_count': '侧边栏文章数量。', 'sidebar_comment_count': '侧边栏评论数量。',
        'article_comment_count': '文章评论分页数量。', 'show_google_adsense': '是否展示广告内容。',
        'google_adsense_codes': '广告嵌入代码，可为空。', 'open_site_comment': '全站评论开关。',
        'color_scheme': '主题枚举；缺省 purple。', 'global_header': '公共头部代码，可为空。',
        'global_footer': '公共尾部代码，可为空。', 'beian_code': '网站备案号，可为空。',
        'analytics_code': '网站统计代码。', 'show_gongan_code': '是否展示公安备案信息。',
        'gongan_beiancode': '公安备案信息，可为空。', 'comment_need_review': '是否需要评论审核。',
    },
    'comments.Comment': {
        'body': '评论正文；模型/表单校验长度 300。', 'author': '评论作者；引用 BlogUser。',
        'article': '评论所属文章；引用 Article。', 'parent_comment': '父评论；为空为顶级评论，仍可有子回复。',
        'is_enable': '是否公开展示；不能单独区分待审与禁用。',
    },
    'comments.CommentReaction': {
        'comment': '目标评论；反向名称 reactions。', 'user': '执行表情互动的用户。',
        'reaction_type': '存储表情字符；选项见枚举说明。', 'created_at': '创建时自动填入当前时间。',
    },
    'oauth.OAuthUser': {
        'author': '可空本地用户关联；为空表示尚未绑定。', 'openid': '第三方身份标识，未单独设唯一。',
        'nickname': '第三方昵称。', 'token': '第三方访问令牌，敏感字段，可为空。',
        'picture': '第三方头像地址，可为空。', 'type': '登录渠道标识，非外键。',
        'email': '第三方邮箱字符串，可为空。', 'metadata': '第三方原始元数据，可为空。',
    },
    'oauth.OAuthConfig': {
        'type': '渠道选项：weibo/google/github/facebook/qq。', 'appkey': '应用标识。',
        'appsecret': '应用密钥，敏感字段。', 'callback_url': 'OAuth 回调地址字符串。',
        'is_enable': '是否启用该渠道配置。',
    },
}


def font_run(run, font='楷体', size=12, bold=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), font)


def body(doc, text, compact=False):
    p = doc.add_paragraph()
    font_run(p.add_run(text), size=10.5 if compact else 12)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.widow_control = True
    if compact:
        snap = OxmlElement('w:snapToGrid')
        snap.set(qn('w:val'), '0')
        p._p.get_or_add_pPr().append(snap)
    return p


def subheading(doc, text, page=False):
    p = doc.add_paragraph()
    font_run(p.add_run(text), font='微软雅黑', size=11, bold=True)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.page_break_before = page
    return p


def table(doc, headers, rows, widths, size=9.5):
    if doc.paragraphs:
        doc.paragraphs[-1].paragraph_format.keep_with_next = True
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    props = t._tbl.tblPr
    layout = props.find(qn('w:tblLayout'))
    if layout is None:
        layout = OxmlElement('w:tblLayout')
        props.append(layout)
    layout.set(qn('w:type'), 'fixed')
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        for key, value in {'val': 'single', 'sz': '4', 'color': 'D0D0D0'}.items():
            el.set(qn('w:' + key), value)
        borders.append(el)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for side, value in [('top', '40'), ('bottom', '40'), ('left', '65'), ('right', '65')]:
        el = OxmlElement('w:' + side)
        el.set(qn('w:w'), value)
        el.set(qn('w:type'), 'dxa')
        margins.append(el)
    props.append(margins)
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    for values in [headers] + list(rows):
        row = t.rows[0] if values is headers else t.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if values is headers:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for cell, value, width in zip(row.cells, values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.02
            p.paragraph_format.widow_control = True
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.keep_with_next = values is headers
            snap = OxmlElement('w:snapToGrid')
            snap.set(qn('w:val'), '0')
            p._p.get_or_add_pPr().append(snap)
            run = p.add_run(str(value))
            font_run(run, '宋体', size, values is headers)
            run._element.rPr.rFonts.set(qn('w:ascii'), 'Times New Roman')
            run._element.rPr.rFonts.set(qn('w:hAnsi'), 'Times New Roman')
            if values is headers:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'), 'F2F2F2')
                cell._tc.get_or_add_tcPr().append(shading)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(1)
    font_run(spacer.add_run(), '宋体', 1)
    snap = OxmlElement('w:snapToGrid')
    snap.set(qn('w:val'), '0')
    spacer._p.get_or_add_pPr().append(snap)
    return t


def format_default(value):
    if value == 'django.utils.timezone.now':
        return 'now'
    if isinstance(value, bool):
        return str(value)
    if value == '':
        return '空字符串'
    return str(value)


def constraints(field):
    items = ['PK，自增'] if field['primary_key'] else ['NULL' if field['null'] else 'NOT NULL']
    if field['unique'] and not field['primary_key']:
        items.append('UQ')
    if field['relation_model']:
        items.append('FK→' + field['relation_model'].split('.')[-1])
        items.append('列=' + field['column'])
        items.append('on_delete=' + field['on_delete'])
    if field['has_default']:
        items.append('默认=' + format_default(field['default']))
    if field['auto_now_add']:
        items.append('auto_now_add')
    if field['auto_now']:
        items.append('auto_now')
    if field['blank']:
        items.append('blank=True')
    if field['choices']:
        items.append('choices 校验')
    return '；'.join(items)


def field_type(field):
    result = field['type']
    if field['max_length']:
        result += '(' + str(field['max_length']) + ')'
    return result + '\n' + str(field['sql_type']).replace(' AUTO_INCREMENT', '')


def field_description(label, field):
    name = field['name']
    if name in DESCRIPTIONS[label]:
        return DESCRIPTIONS[label][name]
    if name in COMMON_FIELDS:
        return COMMON_FIELDS[name]
    raise ValueError('Missing field description: ' + label + '.' + name)


def diagram_font(size, bold=False):
    windows = Path('C:/Windows/Fonts')
    candidates = [windows / ('msyhbd.ttc' if bold else 'msyh.ttc'), windows / 'simhei.ttf']
    for path in candidates:
        if path.exists():
            return ImageFont.truetype(str(path), size)
    raise RuntimeError('ER 图需要可用中文字体；请配置 diagram_font() 的字体路径。')


def draw_diagrams():
    ASSETS.mkdir(parents=True, exist_ok=True)
    title_font = diagram_font(48, True)
    name_font = diagram_font(42, True)
    line_font = diagram_font(37)
    edge_font = diagram_font(34)

    def node(draw, box, name, lines):
        x, y, w, h = box
        draw.rectangle((x, y, x+w, y+h), fill='white', outline='#303030', width=4)
        draw.rectangle((x+2, y+2, x+w-2, y+74), fill='#F2F2F2')
        draw.text((x+24, y+15), name, font=name_font, fill='#151515')
        for n, line in enumerate(lines):
            draw.text((x+24, y+94+n*49), line, font=line_font, fill='#222222')

    def caption(draw, xy, text, anchor='mm'):
        bounds = draw.textbbox(xy, text, font=edge_font, anchor=anchor)
        draw.rectangle((bounds[0]-8, bounds[1]-5, bounds[2]+8, bounds[3]+5), fill='white')
        draw.text(xy, text, font=edge_font, fill='#202020', anchor=anchor)

    def connector(draw, points, label, xy, start, end):
        draw.line(points, fill='#555555', width=4, joint='curve')
        caption(draw, xy, label)
        caption(draw, start[0], start[1])
        caption(draw, end[0], end[1])

    image = Image.new('RGB', (1800, 1480), 'white')
    draw = ImageDraw.Draw(image)
    draw.text((60, 35), '核心内容与互动实体关系', font=title_font, fill='black')
    boxes = {
        'user': (60, 170, 460, 260), 'article': (670, 170, 460, 260), 'category': (1280, 170, 460, 260),
        'reaction': (60, 1040, 460, 260), 'comment': (670, 1040, 460, 260), 'tag': (1280, 1040, 460, 260),
    }
    for key, name, lines in [
        ('user', 'BlogUser 用户', ['PK id', 'UQ username', 'password / email']),
        ('article', 'Article 文章', ['PK id / UQ title', 'FK author_id', 'FK category_id']),
        ('category', 'Category 分类', ['PK id / UQ name', 'FK parent_category_id', 'slug / index']),
        ('reaction', 'CommentReaction', ['PK id / reaction_type', 'FK comment_id', 'FK user_id']),
        ('comment', 'Comment 评论', ['PK id / is_enable', 'FK author_id / article_id', 'FK parent_comment_id']),
        ('tag', 'Tag 标签', ['PK id / UQ name', 'slug', 'Article.tags 关联']),
    ]:
        node(draw, boxes[key], name, lines)
    connector(draw, [(520, 300), (670, 300)], '撰写', (595, 250), ((545, 300), '1'), ((640, 300), '0..N'))
    connector(draw, [(1130, 300), (1280, 300)], '归属', (1205, 250), ((1158, 300), '0..N'), ((1255, 300), '1'))
    connector(draw, [(290, 430), (290, 1040)], '提交表情', (290, 760), ((290, 465), '1'), ((290, 1004), '0..N'))
    connector(draw, [(900, 430), (900, 1040)], '拥有评论', (900, 725), ((900, 470), '1'), ((900, 1004), '0..N'))
    connector(draw, [(455, 430), (455, 630), (730, 835), (730, 1040)], '发表评论', (568, 715), ((455, 477), '1'), ((730, 1004), '0..N'))
    connector(draw, [(670, 1170), (520, 1170)], '评论表情', (595, 1114), ((640, 1170), '1'), ((545, 1170), '0..N'))
    connector(draw, [(1060, 430), (1060, 620), (1510, 845), (1510, 1040)], '标签关联（隐式关联表）', (1280, 740), ((1060, 478), '0..N'), ((1510, 1003), '0..N'))
    draw.line([(1690, 170), (1690, 120), (1780, 120), (1780, 505), (1640, 505), (1640, 430)], fill='#555555', width=4)
    caption(draw, (1500, 117), '父分类 0..1 / 子分类 0..N')
    draw.line([(1130, 1090), (1220, 1090), (1220, 1345), (1035, 1345), (1035, 1300)], fill='#555555', width=4)
    caption(draw, (950, 1393), '父评论 0..1 / 子回复 0..N')
    draw.text((60, 1435), 'PK=主键；FK=外键；UQ=唯一。线端重数表示该端实体对应的记录数。', font=diagram_font(30), fill='#333333')
    image.save(ASSETS / 'er-core.png', dpi=(300, 300))

    image = Image.new('RGB', (1800, 1380), 'white')
    draw = ImageDraw.Draw(image)
    draw.text((60, 35), '第三方身份与独立配置实体', font=title_font, fill='black')
    node(draw, (60, 160, 530, 280), 'BlogUser 用户', ['PK id', '第三方绑定的本地账号', 'AUTH_USER_MODEL'])
    node(draw, (1120, 160, 620, 280), 'OAuthUser 外部身份', ['PK id / type / openid', 'FK author_id（可空）', 'token / metadata（敏感）'])
    connector(draw, [(590, 300), (1120, 300)], '绑定本地账号', (855, 240), ((625, 300), '0..1'), ((1078, 300), '0..N'))
    for box, name, lines in [
        ((60, 605, 740, 275), 'OAuthConfig 渠道配置', ['PK id / type / is_enable', 'appkey / appsecret（敏感）', 'callback_url']),
        ((1000, 605, 740, 275), 'BlogSettings 站点配置', ['PK id / site_name', 'color_scheme / open_site_comment', 'comment_need_review']),
        ((60, 975, 740, 275), 'Links 友情链接', ['PK id / UQ name / UQ sequence', 'link / show_type / is_enable']),
        ((1000, 975, 740, 275), 'SideBar 自定义侧边栏', ['PK id / UQ sequence', 'name / content / is_enable']),
    ]:
        node(draw, box, name, lines)
    draw.text((60, 500), 'OAuthUser.type 与 OAuthConfig.type 是渠道标识；模型之间没有数据库外键。', font=line_font, fill='#333333')
    draw.text((60, 1292), '四个配置实体没有声明到其他业务实体的外键，不虚构连接或数据库单例约束。', font=line_font, fill='#333333')
    image.save(ASSETS / 'er-configuration.png', dpi=(300, 300))


def figure(doc, filename, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
    p.add_run().add_picture(str(ASSETS / filename), width=width)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(p.add_run(caption), '宋体', 10.5)
    p.paragraph_format.space_after = Pt(8)


def preserve_package(reference, generated):
    allowed = {'[Content_Types].xml', 'word/document.xml', 'word/_rels/document.xml.rels', 'docProps/core.xml'}
    original_bytes = reference.read_bytes()
    generated_bytes = generated.read_bytes()
    with zipfile.ZipFile(BytesIO(original_bytes)) as old, zipfile.ZipFile(BytesIO(generated_bytes)) as new:
        items = []
        for info in new.infolist():
            data = old.read(info.filename) if info.filename in old.namelist() and info.filename not in allowed else new.read(info.filename)
            items.append((info, data))
        missing = set(old.namelist()) - set(new.namelist())
        for name in missing:
            items.append((old.getinfo(name), old.read(name)))
    stream = BytesIO()
    with zipfile.ZipFile(stream, 'w', zipfile.ZIP_DEFLATED) as out:
        for info, data in items:
            out.writestr(info, data)
    generated.write_bytes(stream.getvalue())
    with zipfile.ZipFile(reference) as old, zipfile.ZipFile(generated) as new:
        assert all(old.read(name) == new.read(name) for name in old.namelist() if name not in allowed), '模板不允许改动的部件未保留'


def build(models, template):
    by_label = {model['label']: model for model in models}
    if set(by_label) != set(MODEL_NAMES):
        raise ValueError('业务模型清单发生变化，须先更新说明书的实体与字段解释。')
    for model in models:
        for field in model['fields']:
            if field['type'] != field['migration_type']:
                raise ValueError('模型与迁移类型不一致：' + model['label'] + '.' + field['name'])
            field_description(model['label'], field)
    draw_diagrams()
    doc = Document(template)
    reference_paragraphs = [deepcopy(p._p) for p in doc.paragraphs]
    title, group, members = reference_paragraphs[:3]
    sections = [reference_paragraphs[n] for n in (3, 5, 7)]
    for element in list(doc._element.body):
        if element.tag != qn('w:sectPr'):
            doc._element.body.remove(element)

    def original_paragraph(element, text=None, page=False):
        from docx.text.paragraph import Paragraph
        doc._element.body.insert(len(doc._element.body)-1, deepcopy(element))
        p = Paragraph(doc._element.body[-2], doc._body)
        if text is not None:
            original_run = deepcopy(p.runs[0]._r.rPr) if p.runs and p.runs[0]._r.rPr is not None else None
            p.clear()
            r = p.add_run(text)
            if original_run is not None:
                r._r.insert(0, original_run)
        p.paragraph_format.keep_with_next = True
        if page:
            p.paragraph_format.page_break_before = True
        return p

    original_paragraph(title, 'AI Agent 工程实践\n博客系统数据模型设计')
    original_paragraph(group, '组号：24 级软件工程 3 班第 6 组')
    original_paragraph(members, '\n'.join('、'.join(name + '（' + number + '）' for name, number in MEMBERS[i:i+3]) for i in (0, 3)))
    p = body(doc, '版本：1.0（提交版）\n编制：石兆翔　审核责任人：石兆翔\n编制日期：2026 年 10 月 11 日', compact=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    original_paragraph(sections[0])
    subheading(doc, '1.1 系统定位与目标')
    body(doc, '本系统是在 DjangoBlog 基础上二次开发的技术博客，以“AI Agent 工程实践”为内容主题，服务于学习笔记发布、案例检索和读者讨论。主题是文章内容的组织方式，不表示系统已提供自主 Agent 执行、模型训练或智能体编排功能。')
    subheading(doc, '1.2 用户与使用范围')
    body(doc, '访客浏览文章并使用分类、标签、作者、归档和搜索入口；登录用户参与评论、回复和表情互动；具有相应权限的工作人员通过 Django 管理后台维护内容和配置。身份权限沿用 Django 认证体系，后台身份与普通登录用户不能混同。')
    subheading(doc, '1.3 设计依据与边界')
    body(doc, '依据课程第五周要求和《软件数据模型设计说明书模板》，结合 src/DjangoBlog 下 accounts、blog、comments、oauth 模块的模型及其迁移文件编制。分析包含 11 个项目模型、108 个实体字段和 3 个隐式多对多关联表；框架内置权限、会话、日志及插件数据不作为小组新建业务实体重复计算。')
    body(doc, '本说明书描述现有实现，不修改数据库模式、不导出真实业务记录，也不写入密码或第三方凭据。代码核对基于已集成组员注释的 develop 内容；本轮修正位于组长个人分支，最终远程提交仍须经过组内集成验收。', compact=True)

    original_paragraph(sections[1], page=True)
    subheading(doc, '2.1 功能与数据需求')
    table(doc, ['功能', '主要数据', '处理与使用规则'], [
        ('账号与身份', 'BlogUser、Django 权限关联', '以用户名登录；激活、后台访问和超级用户权限分别由相应标志控制。'),
        ('文章发布与浏览', 'Article、BlogUser、Category', '文章关联一名作者和一个分类；status 与 type 区分发布/草稿和文章/独立页面。'),
        ('分类与标签导航', 'Category、Tag、Article.tags', '分类自关联形成层级；标签通过中间表支持跨栏目检索。'),
        ('评论与回复', 'Comment、Article、BlogUser', '记录作者和文章，父评论可空；公开列表过滤 is_enable=True。'),
        ('评论表情互动', 'CommentReaction', '用户针对评论提交表情；三字段联合唯一约束控制同种反应重复。'),
        ('第三方登录', 'OAuthUser、OAuthConfig', '按渠道读取配置和查找外部身份；需要实际配置后才可使用，文档不宣称所有渠道已部署。'),
        ('站点内容配置', 'BlogSettings、Links、SideBar', '维护站点信息、主题、评论策略、链接与侧边栏；不依靠额外的配置外键。'),
        ('搜索与展示', 'Article 及展示上下文', '检索索引由搜索组件维护；搜索结果不是新增的业务实体，也不把缓存当作权威数据库。'),
    ], [1.08, 1.68, 3.0], 10)
    subheading(doc, '2.2 典型数据处理链')
    body(doc, '文章链：后台录入 → 保存 Article 和标签关联 → 列表筛选已发布文章 → 详情读取内容并累加浏览量。')
    body(doc, '评论链：登录用户提交 → 校验并保存 Comment → 按配置设置可见性。表情接口按评论、用户、表情切换反应并统计。')
    body(doc, '第三方身份链：读取渠道配置 → 查询外部身份 → 关联 BlogUser 并登录。查询键不等于数据库唯一键。')
    subheading(doc, '2.3 测试数据与数据质量')
    body(doc, '沿用第四周 AI Agent 主题演示数据方案：12 篇文章、1 个根栏目与 5 个子栏目、14 个标签、3 条可见示例评论。数据生成入口为 blog/management/commands/create_testdata.py；本文不以此宣称当前每台机器已有相同记录数量。演示账号不得视为真实用户或管理员凭据。', compact=True)
    body(doc, '数据质量核对重点：文章作者/分类有效，分类树与回复层级无循环，回复与父评论属于同一文章，反应无重复，敏感字段不进入公开页面或仓库。外键只能确保目标记录存在；“无循环”和“同文章回复”等业务规则不能仅凭外键认定已受数据库保障。', compact=True)

    original_paragraph(sections[2], page=True)
    subheading(doc, '3.1 实体关系图')
    body(doc, '图 1 展示核心内容与互动关系；图 2 展示第三方身份及独立配置实体。重数使用 1、0..1、0..N 表示；线段无箭头，含义是关系而不是页面跳转或数据流方向。Article 与 Tag 的多对多关系由 blog_article_tags 实现。')
    figure(doc, 'er-core.png', '图 1　核心内容与互动 ER 图（关键字段；完整字典见 3.4）')
    body(doc, '父分类和父评论均可空，分别对应根栏目和顶级评论；两者都可以有零个或多个子记录。CommentReaction 同时关联用户与评论，不能只画到评论而遗漏操作者。', compact=True)
    subheading(doc, '3.1.1 外围实体关系', page=True)
    figure(doc, 'er-configuration.png', '图 2　第三方身份与配置实体 ER 图')
    body(doc, 'OAuthUser.author 允许为空，因此一个外部身份对应 0 或 1 个本站用户，一个本站用户可有 0 到多条外部身份记录。OAuthConfig 与 OAuthUser 虽均有 type 字段，但未声明数据库外键；站点配置、链接及侧边栏也是独立实体。', compact=True)
    subheading(doc, '3.2 实体、模型与表的对应关系', page=True)
    table(doc, ['实体 / 模型', '实际表名', '源码位置'], [
        (MODEL_NAMES[m['label']][0] + '\n' + m['label'], m['table'], MODEL_NAMES[m['label']][1]) for m in models
    ], [2.0, 2.15, 1.61], 10)
    body(doc, 'BaseModel 是 abstract=True 的抽象模型，不生成数据表；Article、Category、Tag 继承其 AutoField 主键及两个时间字段。其余模型的隐式主键为 BigAutoField，不能将所有表一概写成 AutoField。', compact=True)
    subheading(doc, '3.3 关系重数与删除行为', page=True)
    relations = [
        ('BlogUser — Article', '1 ↔ 0..N', 'Article.author；CASCADE'),
        ('Category — Article', '1 ↔ 0..N', 'Article.category；CASCADE'),
        ('Article — Tag', '0..N ↔ 0..N', 'Article.tags；隐式关联表'),
        ('BlogUser — Comment', '1 ↔ 0..N', 'Comment.author；CASCADE'),
        ('Article — Comment', '1 ↔ 0..N', 'Comment.article；CASCADE'),
        ('Category（父 — 子）', '0..1 ↔ 0..N', 'Category.parent_category；CASCADE'),
        ('Comment（父 — 子）', '0..1 ↔ 0..N', 'Comment.parent_comment；CASCADE'),
        ('Comment — CommentReaction', '1 ↔ 0..N', 'CommentReaction.comment；CASCADE'),
        ('BlogUser — CommentReaction', '1 ↔ 0..N', 'CommentReaction.user；CASCADE'),
        ('BlogUser — OAuthUser', '0..1 ↔ 0..N', 'OAuthUser.author；可空；CASCADE'),
        ('BlogUser — Group / Permission', '0..N ↔ 0..N', '继承的 groups / user_permissions'),
    ]
    table(doc, ['关系两端', '重数', '实现与删除策略'], relations, [2.03, 1.02, 2.71], 10)
    body(doc, '表中重数按关系两端同一顺序书写。例如 BlogUser—Article 的“1 ↔ 0..N”表示每篇文章必须有一名用户作者，而用户可以尚未发表文章。可空外键不表示级联删除会置空：OAuthUser.author 的 on_delete 仍为 CASCADE。')
    body(doc, 'CASCADE 是 Django 删除收集器采用的级联策略，不应据此承诺直接执行原生 SQL DELETE 也具有完全相同的删除行为。删除作者、文章、分类或父评论可能删除关联内容，实际操作应先确认影响并备份。')

    ordered_groups = [
        ['accounts.BlogUser'], ['blog.Article'], ['blog.Category', 'blog.Tag'],
        ['blog.Links', 'blog.SideBar'], ['blog.BlogSettings'],
        ['comments.Comment'], ['comments.CommentReaction'], ['oauth.OAuthUser'], ['oauth.OAuthConfig'],
    ]
    number = 0
    for group_models in ordered_groups:
        for within, label in enumerate(group_models):
            number += 1
            model = by_label[label]
            subheading(doc, f'3.4.{number} {label}（{MODEL_NAMES[label][0]}）', page=within == 0)
            body(doc, '数据表：' + model['table'] + '。' + MODEL_NAMES[label][2], compact=True)
            if number == 1:
                body(doc, '字段字典图例：PK=主键，UQ=唯一，FK=外键；NULL 描述数据库可空，blank=True 描述校验可留空。类型第二行是当前 MySQL 后端映射；TextField 的括号长度不等于数据库 TEXT 列物理上限。', compact=True)
            table(doc, ['字段', '模型 / SQL 类型', '约束与缺省值', '用途'], [
                (field['name'], field_type(field), constraints(field), field_description(label, field)) for field in model['fields']
            ], [1.43, 1.26, 1.58, 1.49])
            notes = {
                'accounts.BlogUser': '继承的 groups、user_permissions 不作为本表普通字段，见 3.5。RegisterForm 检查邮箱重复，但 email 没有数据库唯一约束。',
                'blog.Article': 'tags 是多对多字段，不存在 Article.tags 普通数据库列。正文 MDTextField 是编辑器提供的文本字段。默认排序为 -article_order、-pub_time。',
                'blog.Category': 'slug 缺省 no-slug；BaseModel.save() 会按 name 生成 slug。父外键不防止分类循环；删除父分类会影响其后代和关联文章。',
                'blog.Tag': 'slug 没有 unique=True；标签名称唯一不表示所有模型的 slug 都唯一。',
                'blog.BlogSettings': '主题枚举：purple、blue、green、orange、pink、red、indigo、teal。clean() 约束只能有一条配置，但普通 save() 不自动调用 full_clean()；未声明数据库单例约束。',
                'comments.Comment': 'is_enable 默认 False；免审核时视图可直接设 True。“公开”不等于已由管理员审核，False 也不能单独区分待审和禁用。默认 now 的修改时间不会自动随更新改变。',
                'comments.CommentReaction': 'reaction_type 存储 👍、👎、❤️、😄、🎉、😕、🚀、👀 中的表情字符；choices 属模型/表单校验，不是数据库枚举 CHECK。联合唯一键见 3.6。',
                'oauth.OAuthUser': 'author 可空，type 与 openid 均未设唯一；业务身份查询不等于联合唯一约束。token、metadata 不应公开展示。',
                'oauth.OAuthConfig': 'type 默认 a 不在声明的渠道 choices 中，是现有实现的配置风险，不能描述为已修复。clean() 检查同 type 的全部记录，不只检查已启用配置。密钥仅说明字段结构，不展示实际值。',
            }
            if label in notes:
                body(doc, notes[label], compact=True)

    subheading(doc, '3.5 隐式多对多关联表', page=True)
    body(doc, '以下三张中间表由 Django 自动创建，不是额外手写的业务模型。各表有自增主键及两个外键，外键组合具有联合唯一约束；既避免重复绑定，又保留一端对多端的关系。', compact=True)
    through_descriptions = {
        'blog_article_tags': {'article': '所关联文章。', 'tag': '所关联标签。'},
        'accounts_bloguser_groups': {'bloguser': '所属用户。', 'group': 'Django 权限组。'},
        'accounts_bloguser_user_permissions': {'bloguser': '所属用户。', 'permission': 'Django 单项权限。'},
    }
    for model in models:
        for field in model['many_to_many']:
            subheading(doc, field['relation_table'])
            body(doc, f"来源：{model['label']}.{field['name']}。联合唯一：" + ', '.join('/'.join(item) for item in field['unique_together']) + '。', compact=True)
            table(doc, ['数据库列', '模型 / SQL 类型', '约束', '含义'], [
                (f['column'], field_type(f), 'PK，自增' if f['primary_key'] else 'NOT NULL；FK→' + f['relation_model'].split('.')[-1] + '；CASCADE', '自增主键。' if f['primary_key'] else through_descriptions[field['relation_table']][f['name']]) for f in field['through_fields']
            ], [1.43, 1.26, 1.58, 1.49])
    body(doc, 'Group 与 Permission 是 Django 认证框架模型。auth_group_permissions 连接组与权限，属于框架支持表；本文不把其字段计入 11 个项目业务模型。会话、站点及管理日志等内置表同理。', compact=True)

    subheading(doc, '3.6 数据库约束与查询索引', page=True)
    body(doc, '数据库唯一约束包括：各表主键；BlogUser.username；Article.title；Category.name；Tag.name；Links.name、Links.sequence；SideBar.sequence；CommentReaction(comment, user, reaction_type)，以及三张多对多中间表的外键组合。它们与 clean()、表单校验和 choices 不是同一层保障。')
    indexes = [(m['label'].split('.')[-1], i['name'], ', '.join(i['fields'])) for m in models for i in m['indexes']]
    table(doc, ['模型', '显式索引名称', '字段顺序'], indexes, [1.05, 2.4, 2.31], 10)
    body(doc, '以上列出 Meta.indexes 中的 7 个显式复合索引；主键、唯一键及外键的后端索引另由 ORM/数据库管理。字段前的“-”表示模型声明降序，实际建表及优化器行为还需结合本机 MySQL 版本与迁移结果核对。', compact=True)
    subheading(doc, '3.7 校验边界与现有实现风险')
    body(doc, '（1）邮箱重复由注册表单检查，不是数据库唯一；直接 ORM 写入以及并发提交不能仅凭表单检查保证唯一。\n（2）BlogSettings.clean() 与 OAuthConfig.clean() 是模型校验；save() 不自动 full_clean()，没有对应数据库唯一键。\n（3）TextField(max_length=300) 与 choices 在校验层发挥作用，不是 TEXT 存储上限或数据库枚举约束。\n（4）分类和评论自关联未在数据库层保证无循环；评论父子是否同文章也应单独核对。\n（5）creation_time、last_modify_time 等 default=now 只提供初值，只有显式 auto_now_add 的 created_at 自动记录创建时间。')
    body(doc, '现有 Article.status 只声明 d/p，CommentPostView 中额外检查 status==c 不能解读为已实现“文章禁用状态”。本文记录这一边界，本轮仅修正错误注释，不修改业务判断或新增数据库约束。若后续修复，应由组内另行确认需求、补测试并生成相应迁移。', compact=True)

    subheading(doc, '3.8 数据安全与维护', page=True)
    body(doc, 'Django 用户密码按认证框架编码保存，不提供通用管理员口令；OAuthUser.token 与 OAuthConfig.appsecret 属敏感字段。文档、图示、代码注释和 Git 提交均不得包含真实密码、令牌或环境凭据。演示使用独立数据库与虚构内容，避免覆盖真实数据。')
    body(doc, '模式变更通过 Django migrations 管理。更新前先同步远程并核对本地改动，在个人分支实施；迁移前备份并验证回滚路径。站点配置保存会清理缓存，但缓存不是数据库实体，搜索索引也需按项目流程独立更新。删除涉及 CASCADE 的对象前必须评估文章、评论、回复及反应的连带影响。')
    subheading(doc, '3.9 文档核对与提交检查')
    table(doc, ['检查项', '本说明书的核对方式'], [
        ('模型与迁移', '只读读取模型元数据，与迁移状态逐字段核对类型；覆盖 11 个实体、108 个字段。'),
        ('关系与约束', '对照 ForeignKey、ManyToManyField、unique、unique_together、indexes、clean()，分层说明真实约束。'),
        ('图文一致', '两张 ER 图片与实体映射、关系重数、字段字典相互对照；不遗漏用户到反应或两个自关联。'),
        ('课程模板', '保留原模板的 A4 页幅、页边距、标题、三节结构、字体及编号；扩展正式内容并逐页渲染检查。'),
        ('资料归档', 'doc/week5 保存正式文件、图示和生成脚本；第四周初稿保留，不覆盖第四周界面文档。'),
        ('最终提交', '个人分支提交 → develop 集成验收 → 所有成员完成后统一合入 main；在远程 main/doc/week5 确认文件。'),
    ], [1.25, 4.51], 10)
    body(doc, '编制者及审核责任人：石兆翔。版本 1.0，编制日期 2026 年 10 月 11 日。本说明书是可提交文档，不代替组员个人贡献、实际签字或远程发布记录；第五周组会信息与会后进度另见《第五周计划与组会记录》。', compact=True)

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    font_run(footer.add_run('第 '), '宋体', 9)
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    font_run(footer.add_run(' 页'), '宋体', 9)
    doc.core_properties.title = 'AI Agent 工程实践博客系统数据模型设计说明书'
    doc.core_properties.author = '石兆翔'
    doc.core_properties.subject = '软件工程方法学第五周数据模型设计'
    doc.core_properties.comments = '课程模板提交版，含模型元数据核对和原生 ER 图。'
    doc.core_properties.modified = datetime(2026, 10, 11, tzinfo=timezone.utc)
    doc.save(OUTPUT)
    preserve_package(template, OUTPUT)
    actual = Document(OUTPUT)
    all_text = '\n'.join(p.text for p in actual.paragraphs) + '\n'.join(c.text for t in actual.tables for row in t.rows for c in row.cells)
    for forbidden in ('xx、', '组号：xx', '待填写', 'TODO', '版本：初稿'):
        if forbidden in all_text:
            raise ValueError('未清理模板占位内容：' + forbidden)
    assert len(actual.inline_shapes) == 2
    assert sum(len(m['fields']) for m in models) == 108
    print(json.dumps({'output': str(OUTPUT), 'models': len(models), 'fields': 108, 'many_to_many_tables': 3, 'images': 2, 'tables': len(actual.tables), 'template_parts_preserved': True}, ensure_ascii=False))


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    parser = argparse.ArgumentParser()
    parser.add_argument('--metadata-stdin', action='store_true', required=True)
    parser.add_argument('--template', type=Path, default=TEMPLATE)
    args = parser.parse_args()
    lines = sys.stdin.read().splitlines()
    metadata = next((line[len('MODEL_METADATA='):] for line in lines if line.startswith('MODEL_METADATA=')), None)
    if metadata is None:
        raise ValueError('未收到 inspect_models.py 输出的 MODEL_METADATA。')
    build(json.loads(metadata), args.template)


if __name__ == '__main__':
    main()
